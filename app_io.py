"""Glue between the Streamlit v5 UI and the ceda-column-design skill (engine/ceda-column-design = the skill folder, verbatim).

The app never re-implements a design rule or an input rule: uploads are read, routed and merged by the skill's
prepare_inputs.py / input_workbook.py, the design is run by run.py + make_outputs.py, the column plan by plan_map.py.
This file only adds UI helpers (figures, file lists, the decisions log, the overload confirmation list).
"""
import os, sys, json, re, io, zipfile, pickle, contextlib, subprocess, datetime, copy

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.join(HERE, 'engine', 'ceda-column-design')
SCRIPTS = os.path.join(SKILL, 'scripts')
if SCRIPTS not in sys.path:
    sys.path.insert(0, SCRIPTS)
import etabs_io as E              # noqa: E402
import layout as LY               # noqa: E402
import design as DS               # noqa: E402
import input_workbook as INPW     # noqa: E402
import prepare_inputs as PI       # noqa: E402
import fc_study as FS             # noqa: E402
import quick_check as QC          # noqa: E402
import fix_tools as FX            # noqa: E402

ENGINE_VERSION = (re.search(r'\(v([\d.]+\w*)\)', open(os.path.join(SKILL, 'SKILL.md'), encoding='utf-8').read()) or [None, '?'])[1]
read_upload, merge, write_tables = PI.read_upload, PI.merge, PI.write_tables
MODELS = (('seismic', 'Seismic model (above ground, reduced stiffness)'),
          ('full_unreduced', 'Full model – Gravity / unreduced stiffness'),
          ('full_reduced', 'Full model – reduced stiffness'))


def empty_tpl():
    return dict(criteria={}, stories=[], columns=[], sections=[], combos_rows=[], combo_cases=[], lap_rows=[], exclude=[],
                approve=[], overload=[], delta_ok=[], ranges=[], prices_rows=[], enlarge=[], splice=[], marks=[], issues=[], missing_keys=[])


# ------------------------------------------------------------------------------------------------ inputs (skill functions)
def prepare(models, combos, tpl):
    """models: {'seismic'|'full_unreduced'|'full_reduced': (tables, units) or None}; combos: (tables, units) of a separate
    Load Combination file or None; tpl: read_template()-style dict (screen + workbook). Skill routing (v1.3.6):
    the unreduced model becomes the primary full model when no reduced-stiffness model is uploaded (ASSUMPTION, reported)."""
    up, route = PI.route_models(models, combos)
    prep = PI.prepare(up, tpl)
    prep['route'] = route
    prep['merge_rep'] = route + prep['merge_rep']
    prep['models_used'] = {k: bool(v) for k, v in up.items()}
    return prep


def input_check(prep):
    """skill input_check without its overload line – overload rows are confirmed on screen after the run (CEDA 21 Oct 2026)."""
    return [c for c in PI.input_check(prep) if not c['input'].startswith('accepted overload')]


def build_cfg(prep, workdir, overload_confirmed, approved=None, in_etabs=False, renumber=False):
    """skill build_cfg; Accept_overload = ONLY the rows the designer confirmed on screen in this session (label, story, dc_max).
    Rows of the input workbook that were not confirmed stay in accept_overload_pending (not applied, reported).
    v5: approved = the ρ approvals applied on the fix board ([] before it – the workbook rows are only pre-selected there);
    in_etabs = round 2+: the enlarged sizes are in the new ETABS export, only the option-2 bar rule stays;
    renumber = Mark FINAL: mark names by section size (locked names dropped)."""
    cfg = PI.build_cfg(prep, workdir)
    if approved is not None:
        smf = cfg['detailing'].get('rho_ask', 0.07) <= 0.06 + 1e-9 and str(cfg['seismic'].get('frame', '')).upper() == 'SMF'
        cfg['approved_rho_over_6'] = [] if smf else [list(x) for x in approved]
    if in_etabs:                       # True = all, or a set of (label, story_from, story_to) already in the ETABS export
        cfg['section_overrides'] = [dict(o, in_etabs=True) if (in_etabs is True or (o['label'], o['story_from'], o['story_to']) in in_etabs)
                                    else o for o in cfg.get('section_overrides', [])]
    if renumber:
        cfg['mark_names'] = {}
    pending = cfg.pop('accept_overload', [])
    cfg['accept_overload'] = [[l, s, float(d)] for (l, s), d in sorted(overload_confirmed.items())]
    left = [r for r in pending if (r[0], r[1]) not in overload_confirmed]
    if left:
        cfg['accept_overload_pending'] = left
    return cfg


# ------------------------------------------------------------------------------------------------ run
def run_engine(workdir, cfg, log=None, documents=False):
    """engine run.py (design only – fast); documents (make_outputs.py --per-mark) only when asked (v4)."""
    import importlib, shutil
    import run as RUN
    importlib.reload(RUN)
    cp = os.path.join(workdir, 'config.json')
    json.dump(cfg, open(cp, 'w'), ensure_ascii=False, indent=1)
    out = os.path.join(workdir, 'OUT')
    if os.path.isdir(out):
        shutil.rmtree(out, ignore_errors=True)       # no stale documents from an earlier run
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        RUN.main(cp, out)
        if documents:
            make_documents(out)
    if log is not None:
        log.append(buf.getvalue())
    return json.load(open(os.path.join(out, 'results.json'))), out


def make_documents(out):
    """calc PDF, back-up Excel (all + per mark), Table of Column PDF/DXF, schedule – the slow part, on request."""
    import importlib
    import make_outputs as MO
    importlib.reload(MO)
    with contextlib.redirect_stdout(io.StringIO()):
        MO.main(out, per_mark=True)
    open(os.path.join(out, 'DOCS_OK'), 'w').write('1')


def documents_ready(out):
    return os.path.exists(os.path.join(out, 'DOCS_OK'))


def fc_whatif(out, prices=None, rule=None):
    """skill fc_study on the finished run (fc_study.json is reused when run.py already made it with the same prices)."""
    p = os.path.join(out, 'fc_study.json')
    if os.path.exists(p) and prices is None and rule is None:
        return json.load(open(p))
    return FS.from_state(out, prices, rule)


def reprice(study, prices):
    """cost of every scenario with new prices, without recomputing the study."""
    for sc in study['scenarios']:
        sc['cost_thb'], sc['cost_missing'] = FS.cost(sc, prices)
    return study


def boq(out, prices=None):
    q = FS.boq(out)
    q['cost_thb'], q['cost_missing'] = FS.cost(q, prices) if prices else (None, [])
    return q


def quick(out, label, story, **kw):
    return QC.check(out, label, story, **kw)


def detail_excel(outdir, mark, bottom_story):
    import make_outputs as MO
    with contextlib.redirect_stdout(io.StringIO()):
        MO.main(outdir, (mark, bottom_story))
    job = json.load(open(os.path.join(outdir, 'results.json')))['config']['project'].get('job', 'job')
    return os.path.join(outdir, f'{job}_{mark}_{bottom_story}_Detailed_Calc.xlsx')


def plan_map(workdir, outdir, story):
    """skill plan_map.py on the primary full model (Point Object Connectivity) -> (pdf, dxf, rows)."""
    job = json.load(open(os.path.join(outdir, 'results.json')))['config']['project'].get('job', 'job')
    pre = os.path.join(outdir, f'{job}_{story}_Column_Mapping')
    r = subprocess.run([sys.executable, '-I', os.path.join(SCRIPTS, 'plan_map.py'), os.path.join(workdir, 'full_reduced.xlsx'),
                        os.path.join(outdir, 'results.json'), pre, story], capture_output=True, text=True)
    if r.returncode:
        raise RuntimeError(r.stderr[-1500:])
    rows = json.load(open(pre + '_rows.json'))
    return pre + '.pdf', pre + '.dxf', rows


def files(outdir):
    res = json.load(open(os.path.join(outdir, 'results.json')))
    job = res['config']['project'].get('job', 'job')
    main = [(f'{job}_Column_Decision.xlsx', 'Decision table (label × story, f\'c study, notes, schedule)'),
            (f'{job}_Column_Decisions_Log.xlsx', 'Decisions log (every designer decision, basis, source, date)'),
            (f'{job}_Column_Calc.pdf', 'Calculation sheets – all marks'),
            (f'{job}_Column_Backup.xlsx', 'Back-up calculation – all marks (live formulas)'),
            (f'{job}_Table_of_Column_DRAFT.dxf', 'Table of Column – DXF (draft)'),
            (f'{job}_Table_of_Column_DRAFT.pdf', 'Table of Column – PDF (draft)'),
            (f'{job}_Column_Schedule_DRAFT.xlsx', 'Schedule with bar coordinates (Revit)'),
            ('results.json', 'Raw engine results')]
    main = [(os.path.join(outdir, f), f, d) for f, d in main if os.path.exists(os.path.join(outdir, f))]
    md = os.path.join(outdir, 'marks')
    per = {}
    if os.path.isdir(md):
        for f in sorted(os.listdir(md)):
            m = re.match(rf'{re.escape(job)}_(.+?)_(Calc\.pdf|Backup\.xlsx)$', f)
            if m:
                per.setdefault(m.group(1), {})[m.group(2)] = os.path.join(md, f)
    return main, per


def zip_all(outdir):
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, 'w', zipfile.ZIP_DEFLATED) as z:
        for root, _, fs in os.walk(outdir):
            for f in fs:
                if f.endswith('.pkl'):
                    continue
                p = os.path.join(root, f)
                z.write(p, os.path.relpath(p, outdir))
    return bio.getvalue()


# ------------------------------------------------------------------------------------------------ review lists
def review_rows(res):
    """label-stories that need the designer: REVISE (D/C > 1 with the most steel that fits), ACCEPTED (overload accepted),
    ASK_RHO (6 % < ρ ≤ 8 %) and δ > 1.4. One row per label-story with what the engine reached."""
    out = []
    for r in res['results']:
        for q in r['per']:
            dmax = max(float(q['dmax_x']), float(q['dmax_y']))          # 'inf' (no capacity) arrives as text
            if q['status'] in ('REVISE', 'ACCEPTED', 'ASK_RHO') or dmax > 1.4:
                out.append(dict(status=q['status'], label=r['label'], story=q['story'], mark=r['mark'],
                                section=f"{r['b'] * 10:.0f}×{r['h'] * 10:.0f}" if r['shape'] != 'circ' else f"Ø{r['b'] * 10:.0f}",
                                fc=q['fc'], bars=f"{r['nbar']}-DB{r['db']}", rho_pct=round(q['rho'] * 100, 2), DC=round(q['dc'], 3),
                                delta=round(dmax, 2), combo=q['gov'].get('combo', '')))
    order = {'REVISE': 0, 'ACCEPTED': 1, 'ASK_RHO': 2, 'OK': 3}
    return sorted(out, key=lambda d: (order.get(d['status'], 9), d['label'], d['story']))


# ------------------------------------------------------------------------------------------------ decisions log
def decisions_log(tpl, cfg, res, prep, overload_meta, src='screen / input workbook', extra=None):
    """Excel log of every designer decision used in the run: item, label, story, value, basis tag, source, date."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    today = datetime.date.today().strftime('%d %b %Y')
    rows = []
    add = lambda item, lab, st, val, basis, source=src, when=today: rows.append([item, lab, st, val, basis, source, when])
    for n in prep.get('route', []):
        add('Models', '', '', n, 'Assumption' if 'ASSUMPTION' in n else 'Input', 'uploads')
    m = cfg['materials']
    if m.get('fc_by_story'):
        fcs = {}
        for s in cfg['story_order']:
            if s in m['fc_by_story']:
                fcs.setdefault(m['fc_by_story'][s], []).append(s)
        for f, ss_ in fcs.items():
            add("Column f'c (fixed)", '', f'{ss_[0]} → {ss_[-1]}', f'{f:.0f} ksc', 'Designer (ETABS model)')
    else:
        add("Column f'c", '', '', f"auto study, typical {m['fc_typical']:.0f} ksc", 'Common practice (CEDA f\'c study)')
    add('Materials', '', '', f"fy {m['fy']:.0f}, fyt {m['fyt']:.0f}, slab f'c {m['fc_slab']:.0f} ksc, cover {cfg['detailing']['cover_mm']:.0f} mm, "
        f"ties DB{cfg['detailing']['tie_db']}", 'Designer')
    if cfg['slender'].get('sway'):
        add('Frame sway', '', '', cfg['slender']['sway'], 'Engineering judgment')
    for o in cfg['slender']['overrides']:
        v = ', '.join(f'{k} {o[k]}' for k in ('Lu_m', 'k') if k in o)
        add('Lu / k override', ', '.join(o['labels']) or 'all', ', '.join(o['stories']), v, 'Engineering judgment')
    for r in tpl.get('sections', []):
        add('Section (designer)', r[0], f"{r[1] or 'all'} → {r[2] or 'all'}", f"{r[3] or 'rect'} {r[4]}×{r[5] or r[4]} mm",
            'Designer – checked with the current ETABS forces, update ETABS and rerun')
    for a, b in cfg.get('approved_rho_over_6', []):
        add('ρ 6–8 % approved', a, b, 'ρ ≤ 8 %', 'Engineering judgment (column size restricted)')
    for a, b, c in cfg.get('accept_delta', []):
        add('δ accepted', a, b, f'δ ≤ {c}', 'Engineering judgment (ACI 6.2.5.3 limit 1.4)')
    for a, b, c in cfg.get('accept_overload', []):
        meta = overload_meta.get((a, b), {})
        add('Overload accepted (confirmed on screen)', a, b, f'D/C ≤ {c}', 'Engineering judgment (ACI 318 needs φSn ≥ U)',
            'screen confirmation' + (f" ({meta.get('by')})" if meta.get('by') else ''), meta.get('at', today))
    for a, b, c in cfg.get('accept_overload_pending', []):
        add('Overload in workbook – NOT confirmed, not applied', a, b, f'D/C ≤ {c}', '–', 'input workbook')
    for c in cfg.get('exclude_combos', []):
        add('Combo excluded', '', '', c, 'Designer')
    for e in extra or []:
        add(*e)
    g = cfg['story_ranges'].get('groups')
    if g:
        add('Story ranges fixed', '', '', '; '.join(f'{a}→{b}' if a != b else a for a, b in g), 'Designer')
    me = cfg['mark_envelope']
    add('Mark envelope', '', '', f"{'on' if me['envelope'] else 'off'} – per label ≤ {me['max_weight'] * 100:.0f} %, project ≤ "
        f"{me.get('project_max_weight', 0) * 100:.0f} %", 'CEDA criterion A')
    for n in res.get('notes', []):
        if n.startswith('marks:'):
            add('Marks result', '', '', n, 'Result')
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Decisions'
    ws.append([f"{cfg['project'].get('job', '')} – column design decisions log (engine v{ENGINE_VERSION})"])
    ws['A1'].font = Font(bold=True, size=12)
    hdr = ['No.', 'item', 'label', 'story', 'decision / value', 'basis', 'source', 'date']
    ws.append(hdr)
    for c in ws[2]:
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='D9E1F2')
    for i, r in enumerate(rows, 1):
        ws.append([i] + r)
        if r[0].startswith('Overload'):
            for c in ws[ws.max_row]:
                c.fill = PatternFill('solid', fgColor='FFF2CC')
    for col, w in zip('ABCDEFGH', (5, 34, 16, 16, 60, 44, 24, 14)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A3'
    bio = io.BytesIO(); wb.save(bio)
    return bio.getvalue()


# ------------------------------------------------------------------------------------------------ views
def load_state(outdir):
    return pickle.load(open(os.path.join(outdir, 'state.pkl'), 'rb'))


def mark_band_eval(outdir, mark, band):
    """engine evaluation of every label-story of one mark-band (same functions as the calc sheets)."""
    res = json.load(open(os.path.join(outdir, 'results.json')))
    st = load_state(outdir)
    cache = DS.Cache()
    fy = st['cfg']['materials'].get('fy', 5000)
    items = []
    for r in res['results']:
        if r['mark'] != mark or tuple(r['band']) != tuple(band):
            continue
        for s in r['stories']:
            c = st['cols'][(r['label'], s)]
            ev = DS.evaluate(cache, st['grids'][c['gkey']], c, st['fc_of'][s], fy, r['db'], tuple(r['counts']))
            items.append(dict(label=r['label'], story=s, r=r, c=c, ev=ev))
    return items


def surface_figures(items):
    """plotly: 3D capped φ-surface + demand points, P–Mx, P–My, Mx–My contour at governing Pu (units t, t-m)."""
    import plotly.graph_objects as go
    gov = max(items, key=lambda q: q['ev']['dcmax'])
    S = gov['ev']['S']
    V = S.hull.points * S._scale()
    I, J, K = S.hull.simplices.T
    f3 = go.Figure()
    f3.add_trace(go.Mesh3d(x=V[:, 1] / 1e3, y=V[:, 2] / 1e3, z=V[:, 0] / 1e3, i=I, j=J, k=K, opacity=0.35, color='#8fb3d9',
                           name='φ-surface', hoverinfo='skip'))
    for q in items:
        mg, P = q['ev']['mg'], q['c']['rows']['P']
        dc = q['ev']['dc']
        txt = [f"{q['label']} {q['story']}<br>{m['combo']}<br>D/C {d:.3f}" for m, d in zip(q['c']['meta'], dc)]
        f3.add_trace(go.Scatter3d(x=mg['x']['Md'] / 1e3, y=mg['y']['Md'] / 1e3, z=P / 1e3, mode='markers',
                                  marker=dict(size=3, color=dc, colorscale='RdYlGn_r', cmin=0, cmax=1.1),
                                  text=txt, hoverinfo='text', name=f"{q['label']} {q['story']}"))
    i = gov['ev']['gov']
    gm = gov['ev']['mg']
    gp = (gm['x']['Md'][i] / 1e3, gm['y']['Md'][i] / 1e3, gov['c']['rows']['P'][i] / 1e3)
    f3.add_trace(go.Scatter3d(x=[0, gp[0]], y=[0, gp[1]], z=[0, gp[2]], mode='lines+markers', line=dict(color='red', width=5),
                              name=f"governing ray D/C {gov['ev']['dcmax']:.3f}"))
    f3.update_layout(scene=dict(xaxis_title='Mx (t-m)', yaxis_title='My (t-m)', zaxis_title='P (t)'), height=620,
                     margin=dict(l=0, r=0, t=30, b=0), legend=dict(font=dict(size=9)))
    figs2 = []
    for ax, lab in (('x', 'Mx'), ('y', 'My')):
        C = S.uniaxial(ax)
        f = go.Figure()
        f.add_trace(go.Scatter(x=C[:, 1] / 1e3, y=C[:, 0] / 1e3, mode='lines', name='φPn–φMn'))
        f.add_trace(go.Scatter(x=-C[:, 1] / 1e3, y=C[:, 0] / 1e3, mode='lines', showlegend=False, line=dict(color='#636efa')))
        for q in items:
            f.add_trace(go.Scatter(x=q['ev']['mg'][ax]['Md'] / 1e3, y=q['c']['rows']['P'] / 1e3, mode='markers', marker=dict(size=5),
                                   name=f"{q['label']} {q['story']}"))
        f.update_layout(xaxis_title=f'{lab} magnified (t-m)', yaxis_title='P (t)', height=420, margin=dict(l=10, r=10, t=30, b=10),
                        title=f'P–{lab} (uniaxial section of the surface)')
        figs2.append(f)
    Pg = min(gov['c']['rows']['P'][i], S.phiPnmax * 0.9999)
    Cc = S.contour(Pg)
    fc = go.Figure()
    if len(Cc):
        fc.add_trace(go.Scatter(x=Cc[:, 0] / 1e3, y=Cc[:, 1] / 1e3, mode='lines', name=f'contour at Pu = {Pg / 1e3:.0f} t'))
    fc.add_trace(go.Scatter(x=[gp[0]], y=[gp[1]], mode='markers', marker=dict(size=10, color='red'), name='governing demand'))
    fc.update_layout(xaxis_title='Mx (t-m)', yaxis_title='My (t-m)', height=420, margin=dict(l=10, r=10, t=30, b=10),
                     title='Mx–My at governing Pu', yaxis_scaleanchor='x')
    return f3, figs2, fc, gov


def dxf_png(path, dpi=200):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import ezdxf
    from ezdxf import bbox
    from ezdxf.addons.drawing import RenderContext, Frontend
    from ezdxf.addons.drawing.matplotlib import MatplotlibBackend
    from ezdxf.addons.drawing.config import Configuration, BackgroundPolicy, ColorPolicy
    doc = ezdxf.readfile(path)
    msp = doc.modelspace()
    ext = bbox.extents(msp)
    w, h = (ext.size.x, ext.size.y) if ext.has_data else (16, 11)
    fig = plt.figure(figsize=(24 * w / max(w, h) + 0.5, 24 * h / max(w, h) + 0.5))
    ax = fig.add_axes([0, 0, 1, 1])
    cfg = Configuration(background_policy=BackgroundPolicy.WHITE, color_policy=ColorPolicy.BLACK)
    Frontend(RenderContext(doc), MatplotlibBackend(ax), config=cfg).draw_layout(msp, finalize=True)
    bio = io.BytesIO()
    fig.savefig(bio, format='png', dpi=dpi, facecolor='white')
    plt.close(fig)
    return bio.getvalue()


# ------------------------------------------------------------------------------------------------ project file (v4)
SLOTS = ('seismic', 'full_unreduced', 'full_reduced', 'combos', 'lap', 'workbook')


def save_project(files, tpl, meta):
    """one .zip: uploaded files + inputs on screen + chosen f'c scenario. Overload confirmations are NOT saved (asked every run)."""
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, 'w', zipfile.ZIP_DEFLATED) as z:
        for slot, lst in files.items():
            for name, data in lst or []:
                z.writestr(f'files/{slot}/{name}', data)
        z.writestr('inputs.json', json.dumps(tpl, ensure_ascii=False, default=str, indent=1))
        z.writestr('meta.json', json.dumps(dict(meta, engine=ENGINE_VERSION, app='v5'), ensure_ascii=False, default=str, indent=1))
    return bio.getvalue()


def load_project(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    files = {k: [] for k in SLOTS}
    for n in z.namelist():
        m = re.match(r'files/([^/]+)/(.+)$', n)
        if m and m.group(1) in files:
            files[m.group(1)].append((m.group(2), z.read(n)))
    tpl = json.loads(z.read('inputs.json')) if 'inputs.json' in z.namelist() else empty_tpl()
    meta = json.loads(z.read('meta.json')) if 'meta.json' in z.namelist() else {}
    base = empty_tpl(); base.update(tpl)
    return files, base, meta


# ------------------------------------------------------------------------------------------------ v5 (fix board, rounds, detail)
def fc_cases(out):
    p = os.path.join(out, 'fc_cases.json')
    return json.load(open(p)) if os.path.exists(p) else None


def fix_groups(res, out, rho_ask):
    """failing label-stories (REVISE) of the run, grouped per label and kind: 'rho' (passes with ρ up to the approval limit)
    or 'enlarge'; δ > 1.4 rows separately. One quick check per failing label-story."""
    g, delta = {}, []
    for r in res['results']:
        for q in r['per']:
            dmax = max(float(q['dmax_x']), float(q['dmax_y']))          # 'inf' (no capacity) arrives as text
            if dmax > 1.4 + 1e-9:
                delta.append(dict(label=r['label'], story=q['story'], delta=round(dmax, 2), dc=round(float(q['dc']), 3)))
            if q['status'] not in ('REVISE', 'ASK_RHO'):
                continue
            a = QC.check(out, r['label'], q['story'], rho_max=rho_ask)
            kind = 'rho' if a['ok'] and a['dc'] <= 1.0 + 1e-9 else 'enlarge'
            key = (r['label'], kind)
            sec = f"Ø{r['b'] * 10:.0f}" if r['shape'] == 'circ' else f"{r['b'] * 10:.0f}×{r['h'] * 10:.0f}"
            e = g.setdefault(key, dict(label=r['label'], kind=kind, mark=r['mark'], section=sec, shape=r['shape'], stories=[], dc=[], rho=[],
                                       ap=[], carried=[]))
            e['stories'].append(q['story']); e['dc'].append(round(float(q['dc']), 3)); e['rho'].append(round(float(q['rho']) * 100, 2))
            e['ap'].append(a)
            if q['status'] == 'ASK_RHO' and QC.check(out, r['label'], q['story'])['status'] == 'OK':
                e.setdefault('carried', []).append(q['story'])      # passes at ρ <= 6 % alone – bars continue from the story above
    order = res['stories']
    rows = sorted(g.values(), key=lambda e: (e['label'], e['kind']))
    for e in rows:
        e['stories'] = sorted(e['stories'], key=order.index)
    return rows, delta


def label_story_eval(outdir, label, story):
    res = json.load(open(os.path.join(outdir, 'results.json')))
    st = load_state(outdir)
    r = next(x for x in res['results'] if x['label'] == label and story in x['stories'])
    c = st['cols'][(label, story)]
    ev = DS.evaluate(DS.Cache(), st['grids'][c['gkey']], c, st['fc_of'][story], st['cfg']['materials'].get('fy', 5000), r['db'], tuple(r['counts']))
    q = next(x for x in r['per'] if x['story'] == story)
    return [dict(label=label, story=story, r=r, c=c, ev=ev, q=q)]


def section_png(outdir, label, story, dpi=150):
    """to-scale section of one label-story: perimeter, hoops / crossties of the CEDA tie pattern, bars (bundles) – as built by
    layout.py (the same functions as the Table of Column)."""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle, Circle
    it = label_story_eval(outdir, label, story)[0]
    st = load_state(outdir)
    r, c = it['r'], it['c']
    g = st['grids'][c['gkey']]
    db, counts = r['db'], tuple(r['counts'])
    b, hh = g.b, g.h
    fig, ax = plt.subplots(figsize=(4.6, 4.6 * min(1.6, max(0.6, hh / b))))
    t = g.cover + g.tie / 2
    if g.shape == 'circ':
        ax.add_patch(Circle((b / 2, b / 2), b / 2, fill=False, lw=1.6, color='#1b2230'))
        ax.add_patch(Circle((b / 2, b / 2), b / 2 - t, fill=False, lw=1.3, color='#b7791f'))
    else:
        ax.add_patch(Rectangle((0, 0), b, hh, fill=False, lw=1.6, color='#1b2230'))
        ax.add_patch(Rectangle((t, t), b - 2 * t, hh - 2 * t, fill=False, lw=1.3, color='#b7791f'))
        T = LY.ties(g, db)
        rr = db / 2
        for ya, yb in T['hoops_h']:
            ax.add_patch(Rectangle((t, ya - rr - g.tie / 2), b - 2 * t, yb - ya + db + g.tie, fill=False, lw=1.1, color='#b7791f'))
        for xa, xb in T['hoops_b']:
            ax.add_patch(Rectangle((xa - rr - g.tie / 2, t), xb - xa + db + g.tie, hh - 2 * t, fill=False, lw=1.1, color='#b7791f'))
        for yc in T['cross_h']:
            ax.plot([t, b - t], [yc + rr + g.tie / 2] * 2, color='#b7791f', lw=1.0)
        for xc in T['cross_b']:
            ax.plot([xc + rr + g.tie / 2] * 2, [t, hh - t], color='#b7791f', lw=1.0)
    for x, y, a in LY.bars_cm(g, db, counts):
        ax.add_patch(Circle((x * 10, y * 10), db / 2, color='#1b2230'))
    ax.set_xlim(-60, b + 60); ax.set_ylim(-60, hh + 60); ax.set_aspect('equal'); ax.axis('off')
    ax.annotate('', (0, -35), (b, -35), arrowprops=dict(arrowstyle='<->', lw=0.8, color='#5a6372'))
    ax.text(b / 2, -55, f'{b:.0f}', ha='center', va='top', fontsize=8, color='#2a3242')
    if g.shape != 'circ':
        ax.annotate('', (-35, 0), (-35, hh), arrowprops=dict(arrowstyle='<->', lw=0.8, color='#5a6372'))
        ax.text(-55, hh / 2, f'{hh:.0f}', ha='right', va='center', rotation=90, fontsize=8, color='#2a3242')
    bio = io.BytesIO(); fig.savefig(bio, format='png', dpi=dpi, bbox_inches='tight', facecolor='white'); plt.close(fig)
    return bio.getvalue(), it


def rename_table(old_marks, new_marks):
    return [dict(label=l, locked=old_marks.get(l, ''), final=new_marks.get(l, '')) for l in sorted(new_marks, key=lambda x: (len(x), x))]
