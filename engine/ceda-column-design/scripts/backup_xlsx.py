"""Calculation back-up workbook: live formulas for everything Excel can do (section, Ise, slenderness, δ, Mu) and
Python values (grey) for the radial D/C on the 3D surface plus the capacity point on the same ray (for checking)."""
import math
import numpy as np
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as CL

F = 'Arial'
f9 = Font(name=F, size=9); fb = Font(name=F, size=9, bold=True); fin = Font(name=F, size=9, color='0000FF')
fgr = Font(name=F, size=9, color='666666', italic=True); fT = Font(name=F, size=12, bold=True)
HDR = PatternFill('solid', fgColor='E8EDF2'); PY = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

BASIS = [
    ('Units', 'kg, cm (section); t, t-m (forces). Blue = input, black = formula, grey = Python value'),
    ('Load combinations', 'กฎกระทรวงกำหนดการออกแบบโครงสร้างอาคาร … พ.ศ. 2566 ข้อ 7 (factors checked against ETABS definitions)'),
    ('φ', 'ข้อ 8 วรรคสอง (QC): tied 0.70, spiral 0.75, flexure/tension 0.90, shear 0.85; transition by εt (ACI 318-19 21.2.2)'),
    ('Pn,max', 'φPn,max = 0.80φPo (tied) / 0.85φPo (spiral), Po = 0.85f\'c(Ag−Ast) + fy·Ast (ACI 318-19 22.4.2.1)'),
    ('Strength', 'Strain compatibility, ACI rectangular block, εcu = 0.003, Es = 2.04e6 ksc, fibre grid 30 × n'),
    ('Biaxial', 'Radial D/C = |demand| / |capacity| on the ray through (Pu, Mux, Muy) on the capped 3D φ-surface (no Bresler)'),
    ('Slenderness', 'Non-sway moment magnifier ACI 6.6.4; δ = Cm/(1−Pu/0.75Pc) ≥ 1, Cm = 1, (EI)eff = (0.2EcIg + EsIse)/(1+βd), Pc without φ'),
    ('M2,min', 'Pu(1.5 + 0.03h) cm before δ when slender (ACI 6.6.4.5.4); Mu = sign(M)·max(|M|, M2,min)·δ'),
    ('βd', '[1.4(DL+SDL) + 0.25·1.7LL] / [1.4(DL+SDL) + 1.7LL] (CEDA Rev.1, conservative)'),
    ('δ ≤ 1.4', 'ACI 318-19 6.2.5.3 (since ACI 318-08 10.10.2.1) – reported for engineer review'),
    ('Bars', 'One grid per column, c/c 125–150 mm; bundles at tie corners first; bundle bars at true positions'),
]


def hdr(ws, row, col, titles, h=30):
    for j, t in enumerate(titles):
        c = ws.cell(row, col + j, t); c.font = fb; c.fill = HDR; c.border = BOX
        c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    ws.row_dimensions[row].height = h


def build(path, project, pages):
    """pages: list of dict(name, title, S, bars, db, b, h, fc, fy, rows [dict label story combo P M3 M2 k Lu betad Mxd Myd dc])"""
    wb = openpyxl.Workbook()
    ws0 = wb.active; ws0.title = 'Basis'
    ws0['A1'] = f"COLUMN CALCULATION BACK-UP – {project.get('name', '')}"; ws0['A1'].font = fT
    for i, (a, b) in enumerate(BASIS, 3):
        ws0.cell(i, 1, a).font = fb; ws0.cell(i, 2, b).font = f9
    ws0.column_dimensions['A'].width = 18; ws0.column_dimensions['B'].width = 125
    sm = wb.create_sheet('Summary')
    hdr(sm, 1, 1, ['Sheet', 'Mark / band', 'Section', "f'c", 'Bars', 'Ast (cm²)', 'ρ', 'φPn,max (t)', 'Max D/C', 'δ max', 'n(δ>1.4)', 'Combos'])
    for pi, pg in enumerate(pages):
        S = pg['S']; nm = pg['name'][:28]
        ws = wb.create_sheet(nm)
        ws['A1'] = pg['title']; ws['A1'].font = fT
        inp = [('b (cm) – x', pg['b']), ('h (cm) – y', pg['h']), ("f'c (ksc)", pg['fc']), ('fy (ksc)', pg['fy']), ('Es (ksc)', 2.04e6),
               ('φ compression', S.phi_c), ('α', S.alpha)]
        ws['A3'] = 'INPUTS'; ws['A3'].font = fb
        for i, (a, v) in enumerate(inp):
            ws.cell(4 + i, 1, a).font = f9; ws.cell(4 + i, 2, float(v)).font = fin
        bars = pg['bars']; B0 = 4; B1 = B0 + len(bars)
        ws.cell(3, 10, 'BAR TABLE (true positions, origin at corner)').font = fb
        hdr(ws, B0, 10, ['Bar', 'x (cm)', 'y (cm)', 'A (cm²)', 'x−b/2', 'y−h/2', 'A·dx²', 'A·dy²'])
        for i, (x, y, a) in enumerate(bars):
            r = B0 + 1 + i
            ws.cell(r, 10, i + 1).font = f9
            ws.cell(r, 11, round(x, 3)).font = fin; ws.cell(r, 12, round(y, 3)).font = fin; ws.cell(r, 13, round(a, 4)).font = fin
            ws.cell(r, 14, f'=K{r}-$B$4/2').font = f9; ws.cell(r, 15, f'=L{r}-$B$5/2').font = f9
            ws.cell(r, 16, f'=M{r}*N{r}^2').font = f9; ws.cell(r, 17, f'=M{r}*O{r}^2').font = f9
        der = [('Ag (cm²)', '=B4*B5' if S.shape == 'rect' else '=PI()*B4^2/4'), ('Ast (cm²)', f'=SUM(M{B0+1}:M{B1})'), ('ρ', '=B13/B12'),
               ('Po (t)', '=(0.85*B6*(B12-B13)+B7*B13)/1000'), ('φPn,max (t)', '=B10*B9*B15'), ('Ec (ksc)', '=15100*SQRT(B6)'),
               ('Ig x (cm⁴)', '=B4*B5^3/12' if S.shape == 'rect' else '=PI()*B4^4/64'),
               ('Ig y (cm⁴)', '=B5*B4^3/12' if S.shape == 'rect' else '=PI()*B4^4/64'),
               ('Ise x (cm⁴)', f'=SUM(Q{B0+1}:Q{B1})'), ('Ise y (cm⁴)', f'=SUM(P{B0+1}:P{B1})'),
               ('r x / r y (cm)', '=0.3*B5' if S.shape == 'rect' else '=0.25*B4'), ('', '=0.3*B4' if S.shape == 'rect' else '=0.25*B4')]
        for i, (a, f) in enumerate(der):
            ws.cell(12 + i, 1, a).font = f9; ws.cell(12 + i, 2, f).font = f9
        ws['B14'].number_format = '0.00%'
        rows = pg['rows']
        T0 = max(B1, 26) + 3
        ws.cell(T0 - 1, 1, 'ALL COMBINATIONS (sorted by D/C). Blue k, Lu, βd per column-story.').font = fb
        th = ['Label', 'Story', 'Combo', 'Pu (t)', 'M3=Mx 1st (t-m)', 'M2=My 1st (t-m)', 'k', 'Lu (cm)', 'βd', 'kLu/rx', 'kLu/ry',
              'Pc x (t)', 'Pc y (t)', 'δx', 'δy', 'Mux design', 'Muy design', 'D/C radial', 'φPn on ray', 'φMnx on ray', 'φMny on ray', 'Pu/φPn']
        hdr(ws, T0, 1, th, 42)
        rows = sorted(rows, key=lambda q: -q['dc'])
        for i, q in enumerate(rows):
            r = T0 + 1 + i
            vals = [q['label'], q['story'], q['combo'], q['P'] / 1000, q['M3'] / 1000, q['M2'] / 1000, q['k'], q['Lu'], q['betad']]
            for j, v in enumerate(vals, 1):
                c = ws.cell(r, j, round(v, 4) if isinstance(v, float) else v); c.font = fin
            ws.cell(r, 10, f'=G{r}*H{r}/$B$22').font = f9
            ws.cell(r, 11, f'=G{r}*H{r}/$B$23').font = f9
            ws.cell(r, 12, f'=PI()^2*(0.2*$B$17*$B$18+$B$8*$B$20)/(1+I{r})/(G{r}*H{r})^2/1000').font = f9
            ws.cell(r, 13, f'=PI()^2*(0.2*$B$17*$B$19+$B$8*$B$21)/(1+I{r})/(G{r}*H{r})^2/1000').font = f9
            ws.cell(r, 14, f'=IF(J{r}>22,IF(D{r}>=0.75*L{r},"NG",MAX(1,1/(1-MAX(D{r},0)/(0.75*L{r})))),1)').font = f9
            ws.cell(r, 15, f'=IF(K{r}>22,IF(D{r}>=0.75*M{r},"NG",MAX(1,1/(1-MAX(D{r},0)/(0.75*M{r})))),1)').font = f9
            ws.cell(r, 16, f'=IF(J{r}>22,IF(E{r}<0,-1,1)*MAX(ABS(E{r}),MAX(D{r},0)*(1.5+0.03*$B$5)/100)*N{r},E{r})').font = f9
            ws.cell(r, 17, f'=IF(K{r}>22,IF(F{r}<0,-1,1)*MAX(ABS(F{r}),MAX(D{r},0)*(1.5+0.03*$B$4)/100)*O{r},F{r})').font = f9
            d = q['dc']
            vv = [d, q['P'] / d / 1000 if d else 0, q['Mxd'] / d / 1000 if d else 0, q['Myd'] / d / 1000 if d else 0]
            for j, v in enumerate(vv):
                c = ws.cell(r, 18 + j, round(float(v), 4 if j == 0 else 2)); c.font = fgr; c.fill = PY
            ws.cell(r, 22, f'=IF(S{r}=0,"",D{r}/S{r})').font = f9
        T1 = T0 + len(rows)
        res = [('Max D/C', f'=MAX(R{T0+1}:R{T1})'), ('δ max', f'=MAX(N{T0+1}:O{T1})'),
               ('n(δ>1.4)', f'=COUNTIF(N{T0+1}:O{T1},">1.4")'), ('Combos', f'=COUNTA(C{T0+1}:C{T1})')]
        ws.cell(3, 5, 'RESULTS').font = fb
        for i, (a, f) in enumerate(res):
            ws.cell(4 + i, 5, a).font = f9; ws.cell(4 + i, 6, f).font = fb
        ws.cell(9, 5, 'Grey = Python (3D surface). Pu/φPn on the ray = D/C (check).').font = fgr
        for j, w in enumerate([22, 12, 32, 9, 11, 11, 5, 8, 7, 8, 8, 10, 10, 7, 7, 11, 11, 9, 10, 10, 10, 8], 1):
            ws.column_dimensions[CL(j)].width = w
        ws.freeze_panes = ws.cell(T0 + 1, 4)
        q = f"'{nm}'!"
        vals = [nm, pg['title'], f"{pg['b']*10:.0f}×{pg['h']*10:.0f}", pg['fc'], f"{len(bars)}-DB{pg['db']}", f'={q}B13', f'={q}B14',
                f'={q}B16', f'={q}F4', f'={q}F5', f'={q}F6', f'={q}F7']
        for j, v in enumerate(vals, 1):
            c = sm.cell(pi + 2, j, v); c.font = f9; c.border = BOX
        sm.cell(pi + 2, 7).number_format = '0.00%'
    for j, w in enumerate([22, 40, 11, 6, 12, 9, 7, 10, 8, 7, 8, 7], 1):
        sm.column_dimensions[CL(j)].width = w
    wb.save(path)
