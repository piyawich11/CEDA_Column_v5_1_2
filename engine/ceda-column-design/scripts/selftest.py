"""Self-test – must print ALL PASS before real use.

usage:  python3 -I <skill>/scripts/selftest.py <skill folder>        (run from an empty scratch folder)
        (a path to the scripts folder itself is accepted too)
Correct values: <skill>/tests/expected.json (source of every value noted there); fixtures: <skill>/tests/fixtures/.
T1 layout rules (grid, ties, ladder, 1 % pattern)          T5 small-face bar count + clear spacing
T2 C1 VOCO end-to-end design (engine baseline)              T8 C1 issued drawing bars check D/C <= 1 (S-04-01-01 R3)
T3 ETABS reader + basement load assembly (hand check)       T4 common story ranges
T6 marks (rotation, ties, envelope)                         T7 Excel input path (forces only + CEDA input workbook)
T9 wider grid (c/c <= 160) for 3-4-bar bundles, circle ties (v1.3.1)   T10 inner bars at tie crossings (v1.3.2)
T11 engineer-accepted overload status (v1.3.3)   T12 headerless ETABS sheet + overload gate (v1.3.4)
T13 column mapping plan (plan_map.py, v1.3.5)   T10b inner bars read from the paattern.dxf fixture
T14 input layer v1.3.6: CEDA TABLE 11.04 lap reader, Story_ranges sheet + coverage check, separate combo file (no double
    factors), unreduced full model promoted when no reduced model is given
T15 f'c study v1.4: S0 / S10 / S25 / recommended on a synthetic matrix, D/C floor with the typical-f'c minimum-steel exemption,
    quantities and cost   T16 quick_check of an unchanged column = the run; enlarged section lowers D/C
T17 v1.5 case table A/B   T18 workbook v1.5 (ρ limits, SMF, Enlarge/Splice/Marks)   T19 enlarged section D/C target, splice, mark lock, proposals
"""
import sys, os, json, subprocess, tempfile, shutil

arg = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
SKILL = os.path.dirname(arg) if os.path.exists(os.path.join(arg, 'run.py')) else arg
HERE = os.path.join(SKILL, 'scripts')
FIX = os.path.join(SKILL, 'tests', 'fixtures')
EXP = json.load(open(os.path.join(SKILL, 'tests', 'expected.json'), encoding='utf-8'))
sys.path.insert(0, HERE)
fails = []


def check(name, cond, info=''):
    print(('PASS ' if cond else 'FAIL ') + name + (f'  [{info}]' if info else ''))
    if not cond:
        fails.append(name)


def run(cfg_path, out):
    r = subprocess.run([sys.executable, '-I', os.path.join(HERE, 'run.py'), cfg_path, out], capture_output=True, text=True)
    if r.returncode:
        print(r.stderr[-2000:])
    r.check_returncode()
    return json.load(open(os.path.join(out, 'results.json')))


import layout as LY      # noqa: E402
import design as DS      # noqa: E402
import etabs_io as E     # noqa: E402

# ---------------------------------------------------------------- T1
x = EXP['T1_layout']
g = LY.Grid(450, 2000, 40, 12, 32)
check(f"T1 grid 2000 face: {x['grid_2000_face_positions']} positions, c/c {x['grid_2000_face_cc_mm']}",
      g.nh == x['grid_2000_face_positions'] and abs(g.sp_h - x['grid_2000_face_cc_mm']) < 0.5, f'{g.nh} {g.sp_h:.1f}')
check(f"T1 grid 450 face: {x['grid_450_face_positions']} bars (small-face rule), no note", g.nb == x['grid_450_face_positions'] and not g.notes)
T = LY.ties(g, 32)
check(f"T1 ties symmetric {x['ties_C1_sets']} closed hoops, no crosstie", T['nset'] == x['ties_C1_sets'] and T['ncross'] == 0
      and T['sup_h'] == x['ties_C1_supported_h'])
check('T1 support pattern m=12 (CEDA sketch)', LY._support_pattern(12) == x['support_pattern_m12'])
mb = {k: DS.min_bar_size('rect', 200, int(k)) for k in x['min_bar_long_side_mm']}
check('T1 min bar by long side', all(mb[k] == v for k, v in x['min_bar_long_side_mm'].items()), str(mb))
st20 = LY.ladder(g, (20, 25, 28, 32))
j = DS.pattern_step(st20, 45 * 200)
check(f"T1 C1 1 % pattern = {x['C1_pattern_1pct'][0]}-DB{x['C1_pattern_1pct'][1]}", [sum(st20[j][1]), st20[j][0]] == x['C1_pattern_1pct'])
st = LY.ladder(g, (25, 28, 32))
n = [sum(s[1]) for s in st]
check('T1 ladder: tie corners first, then other positions', n[:10] == x['ladder_first10_bar_counts'] and st[10][2] == 2, str(n[:12]))

# ---------------------------------------------------------------- T5
x = EXP['T5_small_face']
nf = [LY.face_count(d - 136, 32, face_dim=d)[0] for d in x['face_dims_mm']]
check(f"T5 face bars {x['face_dims_mm']} -> {x['face_bars']}", nf == x['face_bars'], str(nf))
g5 = LY.Grid(250, 500, 40, 12, 32)
st5 = LY.ladder(g5, (16, 20, 25, 28, 32))
check(f"T5 250x500: largest single DB{x['max_single_bar_250x500']} (clear spacing), note", max(s[0] for s in st5) == x['max_single_bar_250x500']
      and any('largest' in m for m in g5.notes), str(sorted({s[0] for s in st5})))
T5 = LY.ties(LY.Grid(540, 540, 40, 12, 20), 20)
check('T5 540 face 3 x DB20: middle bar > 150 clear -> supported (ACI 25.7.2.3)', T5['sup_h'] == x['supported_540_face_DB20'] and T5['ok'])

# ---------------------------------------------------------------- T9 wider grid for 3-4-bar bundles (v1.3.1)
x = EXP['T9_wide_grid']
gn = LY.Grid(900, 900, 40, 12, 32); sn = LY.ladder(gn, (25, 28, 32))
check(f"T9 900 face c/c 125-150: {x['normal_900_face_positions']} bars c/c {x['normal_900_cc_mm']}, max {x['normal_900_max'][0]}-DB32 (2-bar bundles)",
      gn.nh == x['normal_900_face_positions'] and abs(gn.sp_h - x['normal_900_cc_mm']) < 0.1 and [sum(sn[-1][1]), max(sn[-1][1])] == x['normal_900_max'])
gw = LY.Grid(900, 900, 40, 12, 32, (125, 160)); sw = [q for q in LY.ladder(gw, (25, 28, 32)) if max(q[1]) >= 3]
check(f"T9 wider grid c/c <= 160: {x['wide_900_face_positions']} bars c/c {x['wide_900_cc_mm']}, 3-4-bar bundles {x['wide_900_first'][0]} -> {x['wide_900_max'][0]}-DB32",
      gw.nh == x['wide_900_face_positions'] and abs(gw.sp_h - x['wide_900_cc_mm']) < 0.1 and [sum(sw[0][1]), max(sw[0][1])] == x['wide_900_first']
      and [sum(sw[-1][1]), max(sw[-1][1])] == x['wide_900_max'])
check('T9 clear spacing DB32 bundles of 3 / 4 (ACI 25.2.3, 25.6.1.6)', abs(LY.clear_req(32, 3) - x['clear_req_DB32_3']) < 0.05
      and abs(LY.clear_req(32, 4) - x['clear_req_DB32_4']) < 0.05)
gc = LY.Grid(900, 900, 40, 12, 32, shape='circ'); gcw = LY.Grid(900, 900, 40, 12, 32, (125, 160), shape='circ')
check('T9 circle Ø900 positions (normal / wider) and ties carry ncross', gc.n == x['circ_900_positions_normal'] and gcw.n == x['circ_900_positions_wide']
      and LY.ties(gc, 32).get('ncross') == 0)

# ---------------------------------------------------------------- T10 inner bars at the tie crossings (v1.3.2)
x = EXP['T10_inner_bars']
for (bb, hh), key, nb_ in (((1100, 900), 'inner_1100x900', 'dxf_pattern_1100x900_bars'), ((900, 900), 'inner_900x900', 'dxf_pattern_900x900_bars')):
    gi = LY.Grid(bb, hh, 40, 12, 32); gi.interior = True
    got = sorted([round(p['x'], 1), round(p['y'], 1)] for p in gi.inner(32))
    Pp, _ = gi._perim(32)
    # rebuild the drawing's arrangement: perimeter (2-bar bundles at the drawn positions) + inner singles
    check(f"T10 {bb}x{hh}: inner bars at the inner-hoop crossings = paattern.dxf", got == sorted(x[key]), str(got))
    cnt = [1] * len(Pp) + [1] * len(got)
    check(f"T10 {bb}x{hh}: positions() = perimeter + inner, bars_cm counts them", len(gi.positions(32)[0]) == len(cnt)
          and len(LY.bars_cm(gi, 32, tuple(cnt))) == len(cnt))
check('T10 describe: inner bars named', 'INNER BARS' in LY.describe((2, 1, 1), 3))
# T10b – the source drawing itself: inner bars read from tests/fixtures/paattern.dxf (circles strictly inside the inner-hoop crossings)
try:
    import ezdxf
    dd = ezdxf.readfile(os.path.join(FIX, 'paattern.dxf')); mm = dd.modelspace()
    outl = [l for l in mm.query('LINE') if l.dxf.layer == 'CED_0_18_Solid']
    ok_all = True
    for (bb, hh), key in (((1100, 900), 'inner_1100x900'), ((900, 900), 'inner_900x900')):
        # concrete outline of this section: the CED_0_18 square of size bb × hh
        xs = sorted({round(min(l.dxf.start.x, l.dxf.end.x)) for l in outl}); found = None
        for x0 in xs:
            seg = [l for l in outl if round(min(l.dxf.start.x, l.dxf.end.x)) >= x0 - 1 and round(max(l.dxf.start.x, l.dxf.end.x)) <= x0 + bb + 1]
            ys = [min(l.dxf.start.y, l.dxf.end.y) for l in seg]
            if seg and abs(max(max(l.dxf.start.x, l.dxf.end.x) for l in seg) - x0 - bb) < 2 and ys:
                y0 = min(ys)
                if abs(max(max(l.dxf.start.y, l.dxf.end.y) for l in seg) - y0 - hh) < 2:
                    found = (x0, y0); break
        x0, y0 = found
        e = 68.0
        cin = sorted([round(c.dxf.center.x - x0, 1), round(c.dxf.center.y - y0, 1)] for c in mm.query('CIRCLE')
                     if x0 + e + 40 < c.dxf.center.x < x0 + bb - e - 40 and y0 + e + 40 < c.dxf.center.y < y0 + hh - e - 40)
        ok_all = ok_all and all(abs(a[0] - b[0]) < 1.0 and abs(a[1] - b[1]) < 1.0 for a, b in zip(cin, sorted(x[key]))) and len(cin) == len(x[key])
    check('T10b inner bars of paattern.dxf (fixture) = engine positions (1100×900: 8, 900×900: 4)', ok_all)
except ImportError:
    check('T10b ezdxf installed (needed to read paattern.dxf)', False)

# ---------------------------------------------------------------- T11 engineer-accepted overload (v1.3.3)
import run as RUN   # noqa: E402
check('T11 status: D/C 1.05 within the accepted 1.10 -> ACCEPTED; 1.12 -> REVISE; back to OK when D/C <= 1',
      RUN._status('OK', 1.05, 1.10) == 'ACCEPTED' and RUN._status('OK', 1.12, 1.10) == 'REVISE'
      and RUN._status('ACCEPTED', 0.98, 1.10) == 'OK' and RUN._status('OK', 1.05, 1.0) == 'REVISE' and RUN._status('ASK_RHO', 0.9, 1.0) == 'ASK_RHO')

tmp = tempfile.mkdtemp()
try:
    # ------------------------------------------------------------ T2
    x = EXP['T2_C1_VOCO']
    for f in ('c1_forces.json', 'c1_config.json'):
        shutil.copy(os.path.join(FIX, f), tmp)
    o = run(os.path.join(tmp, 'c1_config.json'), os.path.join(tmp, 'out'))
    got = {q['story']: (r['nbar'], r['db'], q['dc']) for r in o['results'] for q in r['per']}
    for k, v in x['bands'].items():
        gv = got.get(k)
        check(f'T2 C1 {k} {v[0]}-DB{v[1]} D/C {v[2]}', gv is not None and gv[:2] == (v[0], v[1]) and abs(gv[2] - v[2]) <= x['dc_tol'],
              f'{gv[0]}-DB{gv[1]} {gv[2]:.3f}' if gv else 'missing')
    dy = {q['story']: q['dmax_y'] for r in o['results'] for q in r['per']}
    check(f"T2 C1 1st-3th δy max {x['delta_y_1st_3th']}", abs(dy['1st-3th'] - x['delta_y_1st_3th']) < 0.01, f"{dy['1st-3th']:.3f}")
    check('T2 C1 every story range kept as a row', [list(b) for b in o['bands']] == x['rows'], str(o['bands']))
    # ------------------------------------------------------------ T8 issued drawing
    import pickle
    x = EXP['T8_C1_issued_drawing']
    stt = pickle.load(open(os.path.join(tmp, 'out', 'state.pkl'), 'rb'))
    cache = DS.Cache()
    for band, (nb_, db_, dce) in x['bands'].items():
        gd = LY.Grid(450, 2000, 40, 12, db_); gd.nb, gd.nh, gd.notes = x['drawn_grid']['nb'], x['drawn_grid']['nh'], []
        sd = min([s for s in LY.ladder(gd, (db_,), 4) if sum(s[1]) >= nb_], key=lambda s: sum(s[1]))
        ev = DS.evaluate(cache, gd, stt['cols'][('C1', band)], stt['fc_of'][band], 5000, db_, sd[1])
        check(f'T8 issued C1 {band} {nb_}-DB{db_}: D/C {dce} <= 1', abs(ev['dcmax'] - dce) <= x['dc_tol'] and ev['dcmax'] <= 1.0, f"{ev['dcmax']:.3f}")
    # ------------------------------------------------------------ T6 marks
    def _t6(envelope):
        sts = ['B3', 'B2', 'B1'] + [f'{i}F' for i in range(1, 13)]
        rows = []

        def add(lab, scale, top, V=0.0):
            nn = sts.index(top) + 1
            for i, s_ in enumerate(sts[:nn]):
                P = (nn - i) * 95000 * scale
                for c, (fp, fm) in {'G1': (1, 1), 'E1': (0.75, 1), 'E2': (0.6, -1)}.items():
                    rows.append(dict(label=lab, story=s_, combo=c, P=P * fp, M3=0.04 * P * fm, M2=0.03 * P * fm, V2=V, V3=V / 3,
                                     kind='seismic' if c != 'G1' else 'gravity'))
        add('A1', 1.0, '12F'); add('A2', 1.0, '12F'); add('A3', 1.0, '12F', V=60000)
        add('B1', 1.0, '12F'); add('B2', 1.08, '12F'); add('D1', 1.0, '8F')
        json.dump(dict(rows=rows, stories=sts), open(os.path.join(tmp, 'f6.json'), 'w'))
        sec = {k: dict(b=40, h=80) for k in ('A1', 'A2', 'A3', 'B2', 'D1')}; sec['B1'] = dict(b=80, h=40)
        fc = {s_: (560 if s_ in ('B3', 'B2', 'B1', '1F', '2F', '3F', '4F') else 400) for s_ in sts}
        cfg = dict(project=dict(name='T6', job='T6'), seismic=dict(frame='gravity'), materials=dict(fy=5000, fyt=4000, fc_typical=400, fc_by_story=fc),
                   detailing=dict(cover_mm=40, tie_db=12, kmax=4, so_mm=150, s_mm=300), slender=dict(k=1.0, beta_d=0.6, Lu_default_m=3.0),
                   sections=sec, lap_table={'400': {'25': 1550, '28': 1875, '32': 2325}, '560': {'25': 1325, '28': 1600, '32': 1975}},
                   inputs=dict(forces_json='f6.json'), basement_stories=['B3', 'B2', 'B1'], story_ranges=dict(mode='auto'),
                   mark_envelope=dict(envelope=envelope))
        cp = os.path.join(tmp, f'c6_{envelope}.json'); json.dump(cfg, open(cp, 'w'))
        return run(cp, os.path.join(tmp, f'o6_{envelope}'))
    mk = _t6(False)['marks']
    check('T6 exact: A1 = A2; A3 (other ties) separate; D1 separate', mk['A1'] == mk['A2'] and mk['A3'] != mk['A1']
          and mk['D1'] not in (mk['A1'], mk['A3']), str(mk))
    o6 = _t6(True); mk = o6['marks']
    check('T6 envelope: A1 = A2 = A3; rotated B1 shares a mark; D1 alone', len({mk[k] for k in ('A1', 'A2', 'A3')}) == 1
          and list(mk.values()).count(mk['B1']) > 1 and list(mk.values()).count(mk['D1']) == 1, str(mk))
    sA = {(r['label'], r['band'][0]): (r['s'], r['nbar'], r['db']) for r in o6['results']}
    check('T6 envelope: ties of the mark = critical (A1 takes s 150 of A3)', sA[('A1', '6F')][0] == 150 and sA[('A3', '6F')][0] == 150)
    mate = next(k for k in mk if k != 'B1' and mk[k] == mk['B1'])
    check('T6 envelope: rotated members carry the same bars in every row', all(sA[(mate, b[0])][1:] == sA[('B1', b[0])][1:] for b in o6['bands']))
    check(f"T6 envelope: steel increase <= {EXP['T6_marks']['max_weight_pct']} %",
          all(int(m.split('(+')[1].split(' %')[0]) <= EXP['T6_marks']['max_weight_pct'] for m in o6['notes'] if m.startswith('mark envelope')))
    # ------------------------------------------------------------ T4
    sts = ['B3', 'B2', 'B1'] + [f'{i}F' for i in range(1, 13)]
    fcm = {s_: (560 if s_ in ('B3', 'B2', 'B1', '1F', '2F', '3F', '4F') else 400) for s_ in sts}
    hm = {s_: 3.0 for s_ in sts}; hm['1F'] = 4.5
    stp = [(20, (1,) * 30), (25, (1,) * 30), (28, (1,) * 30), (32, (1,) * 30), (32, (2,) * 4 + (1,) * 26), (32, (2,) * 8 + (1,) * 22)]
    rA = dict(zip(sts, [5, 5, 4, 4, 3, 3, 3, 2, 2, 2, 2, 1, 1, 0, 0])); rB = dict(zip(sts, [3, 3, 3, 2, 2, 2, 2, 1, 1, 1, 1, 1, 0, 0, 0]))
    rg = DS.story_ranges(sts, {'B3', 'B2', 'B1'}, hm, fcm, {'A': rA, 'B': rB}, {'A': stp, 'B': stp}, 2.325)
    check('T4 common story ranges', [list(r) for r in rg] == EXP['T4_story_ranges']['ranges'], str(rg))
    # ------------------------------------------------------------ T3
    x = EXP['T3_etabs_basement']
    subprocess.run([sys.executable, '-I', os.path.join(FIX, 'make_synthetic_etabs.py'), tmp], check=True)
    L = lambda p: dict(zip(('tables', 'units'), E.read_tables(os.path.join(tmp, p))))
    S, FR, FU = L('seis.xlsx'), L('full_red.xlsx'), L('full_unred.xlsx')
    rows, rep = E.assemble(S, FR, FU)
    check('T3 story order / basement', rep['stories_full'] == x['stories'] and rep['basement_stories'] == x['basement'])
    chk = {c['combo']: c['ok'] for c in rep['combo_check_full']}
    check('T3 combo check ข้อ 7', chk == x['combo_ok'], str(chk))
    b = {r['story']: r for r in rows if r['src'] == 'basement' and r['label'] == 'C1'}
    check('T3 basement P (B1, B2)', all(abs(b[s_]['P'] - v) < 1 for s_, v in x['C1_basement_P_kg'].items()), f"{b['B1']['P']:.0f} {b['B2']['P']:.0f}")
    check('T3 basement M3 / V2 carried down', abs(b['B2']['M3'] - x['C1_B2_M3_kgm']) < 1 and abs(b['B2']['V2'] - x['C1_B2_V2_kg']) < 1)
    check('T3 βd B2', abs(E.beta_d(FR, 'C1', 'B2') - x['C1_B2_beta_d']) < 1e-4)
    # ------------------------------------------------------------ T7 Excel input path
    x = EXP['T7_excel_input']
    pi = os.path.join(HERE, 'prepare_inputs.py')
    d0 = os.path.join(tmp, 'p0')
    r0 = subprocess.run([sys.executable, '-I', pi, d0, '--seismic', os.path.join(FIX, 'voco_column_load.xlsx'), '--template'],
                        capture_output=True, text=True)
    miss = [c['input'] for c in json.load(open(os.path.join(d0, 'input_check.json')))['check'] if not c['ok']]
    check('T7 forces only: missing inputs reported (exit 2) + prefilled workbook written', r0.returncode == 2 and miss == x['missing_without_workbook']
          and os.path.exists(os.path.join(d0, 'CEDA_Column_Input_prefilled.xlsx')), str(miss))
    d1 = os.path.join(tmp, 'p1')
    r1 = subprocess.run([sys.executable, '-I', pi, d1, '--seismic', os.path.join(FIX, 'voco_column_load.xlsx'),
                         '--workbook', os.path.join(FIX, 'voco_input_workbook.xlsx')], capture_output=True, text=True)
    check('T7 with the input workbook: every input complete (exit 0)', r1.returncode == 0, r1.stdout[-300:])
    o7 = run(os.path.join(d1, 'config.json'), os.path.join(d1, 'OUT'))
    per = [q for r in o7['results'] for q in r['per']]
    check("T7 f'c / k / βd / Lu from the workbook", o7['fc_of']['3rd'] == x['fc_3rd'] and o7['fc_of']['4th'] == x['fc_4th']
          and all(abs(q['k'] - x['k']) < 1e-9 and abs(q['betad'] - x['beta_d']) < 1e-9 and abs(q['Lu'] - x['Lu_cm']) < 1e-6 for q in per))
    check('T7 story order from the workbook', o7['stories'][0] == x['first_story'] and o7['stories'][-1] == x['last_story'])
    check(f"T7 combo check flags {x['combo_flag']} (factors do not match the name)", any(x['combo_flag'] in m and 'do not match' in m for m in o7['notes']))
    # ------------------------------------------------------------ T12 headerless ETABS sheet + overload gate (v1.3.4)
    import openpyxl
    hp = os.path.join(tmp, 'pts.xlsx'); wbh = openpyxl.Workbook(); wsh = wbh.active; wsh.title = 'Point Object Connectivity'
    for row in (['UniqueName', 'Is Auto Point', 'Story', 'PointBay', 'IsSpecial', 'X', 'Y', 'Z'], [None, None, None, None, None, 'm', 'm', 'm'],
                [11, 'No', 'B2', 1, 'No', 1.5, 2.5, -6.0]):
        wsh.append(row)
    wbh.save(hp)
    Th, Uh = E.read_tables(hp)
    ph = Th.get('Point Object Connectivity', [{}])
    check('T12 sheet without its TABLE: row read by the sheet name (Point Object Connectivity, units row)',
          len(ph) == 1 and abs(float(ph[0].get('X', 0)) - 1.5) < 1e-9 and Uh['Point Object Connectivity'].get('X') == 'm')
    wbo = openpyxl.load_workbook(os.path.join(FIX, 'voco_input_workbook.xlsx'))
    wso = wbo.create_sheet('Accept_overload'); wso.append(['note']); wso.append(['label', 'story', 'dc_max']); wso.append(['C1', 'B6', 1.05])
    wbp = os.path.join(tmp, 'wb_ovl.xlsx'); wbo.save(wbp)
    d2 = os.path.join(tmp, 'p2'); d3 = os.path.join(tmp, 'p3')
    r2 = subprocess.run([sys.executable, '-I', pi, d2, '--seismic', os.path.join(FIX, 'voco_column_load.xlsx'), '--workbook', wbp], capture_output=True, text=True)
    c2 = json.load(open(os.path.join(d2, 'config.json')))
    r3 = subprocess.run([sys.executable, '-I', pi, d3, '--seismic', os.path.join(FIX, 'voco_column_load.xlsx'), '--workbook', wbp,
                         '--overload-confirmed'], capture_output=True, text=True)
    c3 = json.load(open(os.path.join(d3, 'config.json')))
    check('T12 overload gate: Accept_overload not applied until confirmed in the session (exit 2, pending) – CEDA 21 Oct 2026',
          r2.returncode == 2 and not c2.get('accept_overload') and c2.get('accept_overload_pending') and r3.returncode == 0
          and c3.get('accept_overload') == [['C1', 'B6', 1.05]], f'{r2.returncode} {r3.returncode}')
    # ------------------------------------------------------------ T13 column mapping plan (plan_map.py, v1.3.5)
    st0 = o7['stories'][0]
    labs = sorted({r['label'] for r in o7['results']})
    wbm = openpyxl.Workbook(); w1 = wbm.active; w1.title = 'Column Object Connectivity'
    w1.append(['TABLE:  Column Object Connectivity']); w1.append(['Unique Name', 'Story', 'ColumnBay', 'UniquePtI', 'UniquePtJ', 'Length'])
    w1.append(['', '', '', '', '', 'm'])
    w2 = wbm.create_sheet('Point Object Connectivity')
    w2.append(['UniqueName', 'Story', 'X', 'Y', 'Z']); w2.append([None, None, 'm', 'm', 'm'])
    for i, lb in enumerate(labs):
        w1.append([f'U{i}', st0, lb, f'{1000 + i}', f'{2000 + i}', 3.0])
        w2.append([1000 + i, st0, 6.0 * i, 0.0, 0.0]); w2.append([2000 + i, st0, 6.0 * i, 0.0, 3.0])
    mp = os.path.join(tmp, 'map.xlsx'); wbm.save(mp)
    rm = subprocess.run([sys.executable, '-I', os.path.join(HERE, 'plan_map.py'), mp, os.path.join(d1, 'OUT', 'results.json'),
                         os.path.join(tmp, 'map'), st0], capture_output=True, text=True)
    rows_m = json.load(open(os.path.join(tmp, 'map_rows.json'))) if rm.returncode == 0 else []
    check(f'T13 plan_map: {len(labs)} column(s) at {st0} placed with mark + label, PDF + DXF written',
          rm.returncode == 0 and len(rows_m) == len(labs) and all(q['x'] is not None and q['mark'] for q in rows_m)
          and os.path.exists(os.path.join(tmp, 'map.pdf')) and os.path.exists(os.path.join(tmp, 'map.dxf')), rm.stderr[-300:])
finally:
    shutil.rmtree(tmp, ignore_errors=True)

# ---------------------------------------------------------------- T14 input layer (v1.3.6)
import input_workbook as INPW     # noqa: E402
import prepare_inputs as PI      # noqa: E402
x = EXP['T14_input_layer']
rows, _ = INPW.read_lap_table(os.path.join(FIX, 'lap_table_TABLE_11_04.xlsx'), '40mm')
lap = {r[0]: dict(zip((12, 16, 20, 25, 28, 32), r[1:])) for r in rows}
check(f"T14 lap TABLE 11.04 (40mm, Lst): {len(rows)} f'c rows, 400 ksc DB32 {x['lap_400_DB32']}, 600 ksc DB25 {x['lap_600_DB25']}",
      len(rows) == x['lap_rows'] and lap[400][32] == x['lap_400_DB32'] and lap[600][25] == x['lap_600_DB25'] and lap[210][12] == x['lap_210_DB12'])
od = ['B2', 'B1', '+00', 'GF', 'GM']
check('T14 Story_ranges coverage check', INPW.check_ranges([['B2', 'B2'], ['B1', 'B1'], ['+00', 'GF'], ['GM', 'GM']], od) == []
      and len(INPW.check_ranges([['B2', 'B1'], ['GF', 'GM']], od)) == 1)
tmp = tempfile.mkdtemp()
tp = os.path.join(tmp, 'wb.xlsx')
open(tp, 'wb').write(INPW.make_template({'stories': od}, dict(ranges=[['B2', 'B1'], ['+00', 'GM']])))
cfgr = INPW.to_cfg(INPW.read_template(tp))
check('T14 Story_ranges sheet round trip -> story_ranges.groups', cfgr['story_ranges']['groups'] == [['B2', 'B1'], ['+00', 'GM']])
cmb = PI.read_upload(os.path.join(FIX, 'bkmph02_load_combination.xlsx'))
model = ({'Load Combination Definitions': [{'Name': 'U_01_03_0.9 DL + 1.0 EEX2', 'Load Name': 'DL', 'SF': 0.9},
                                           {'Name': None, 'Load Name': 'EEX-2', 'SF': 1.0}]}, {})
up, notes = PI.route_models({'seismic': model, 'full_reduced': None, 'full_unreduced': model}, (cmb[0], cmb[1]))
fac, _ = E.combo_factors(up['full_reduced'][0])
f = fac['U_01_03_0.9 DL + 1.0 EEX2']
check(f"T14 separate combo file replaces per combo (0.9 DL, not 1.8) and unreduced model promoted",
      abs(f['DL'] - 0.9) < 1e-9 and abs(f['SDL'] - 0.9) < 1e-9 and up['full_unreduced'] is None and up['full_reduced'] is not None
      and any('ASSUMPTION' in n for n in notes), str(f))
shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- T15 f'c study v1.4 (CEDA 23 Oct 2026)
import fc_study as FS     # noqa: E402
x = EXP['T15_fc_study']
sts = ['B1', '1F', '2F']
labs = [f'L{i}' for i in range(10)]
cols15 = {(l, s_): dict(shape='rect', b=50.0, h=50.0, H=300.0) for l in labs for s_ in sts}
def _r(ok, dc, rho, at_min=False):
    return dict(ok=ok, dc=dc, rho=rho, As=20.0, db=25, at_min=at_min)
M = {}
for i, l in enumerate(labs):
    M[(l, 'B1')] = {400: _r(i >= 2, 0.9, 0.03), 500: _r(i >= 1, 0.7, 0.02), 600: _r(True, 0.5, 0.01)}
    M[(l, '1F')] = {400: _r(True, 0.8, 0.012 if i < 6 else 0.025), 500: _r(True, 0.6, 0.01), 600: _r(True, 0.5, 0.01)}
    M[(l, '2F')] = {400: _r(True, 0.5, 0.01, True), 500: _r(True, 0.4, 0.01), 600: _r(True, 0.3, 0.01)}
o15 = FS.study(M, cols15, sts, 400, [500, 600], None, dict(concrete={'400': 2600, '500': 2850, '600': 3150}, steel={'25': 23.5}))
nm = lambda i: [o15['scenarios'][i]['fc_bottom'], o15['scenarios'][i]['reduce_at']]
check(f"T15 S0 (no failing) = {x['S0']}, S10 = {x['S10']}, S25 = {x['S25']}, recommended = {x['recommended']}",
      nm(o15['s0']) == x['S0'] and nm(o15['s10']) == x['S10'] and nm(o15['s25']) == x['S25'] and nm(o15['recommended']) == x['recommended'],
      str([nm(o15[k]) for k in ('s0', 's10', 's25', 'recommended')]))
s10 = o15['scenarios'][o15['s10']]
check(f"T15 S10: {x['S10_fail']} failing, {x['S10_under']} under-used (D/C < 0.75); typical f'c + minimum steel exempt",
      s10['n_fail'] == x['S10_fail'] and s10['n_under'] == x['S10_under'], f"{s10['n_fail']} {s10['n_under']}")
sr = o15['scenarios'][o15['recommended']]
check(f"T15 quantities and cost of the recommended scenario: {x['rec_conc_m3']} m³, {x['rec_steel_kg']} kg, THB {x['rec_cost']}",
      abs(sum(sr['concrete_m3'].values()) - x['rec_conc_m3']) < 1e-6 and abs(sum(sr['steel_kg'].values()) - x['rec_steel_kg']) < 0.2
      and sr['cost_thb'] == x['rec_cost'], f"{sum(sr['concrete_m3'].values())} {sum(sr['steel_kg'].values())} {sr['cost_thb']}")

Mf = {k: ({f: dict(v, ok=(f == 600)) for f, v in d.items()} if k[1] == '2F' else d) for k, d in M.items()}   # top story always typical
of = FS.study(Mf, cols15, sts, 400, [500, 600])
check('T15 no scenario within the limits -> FALLBACK recommended (never None), listed for the designer',
      of['recommended'] is not None and any('FALLBACK' in t for t in of['scenarios'][of['recommended']]['tags']) and of['whatif'])

# ---------------------------------------------------------------- T16 quick_check (fix-board preview)
import quick_check as QC  # noqa: E402
tmp16 = tempfile.mkdtemp()
try:
    for f in ('c1_forces.json', 'c1_config.json'):
        shutil.copy(os.path.join(FIX, f), tmp16)
    o16 = run(os.path.join(tmp16, 'c1_config.json'), os.path.join(tmp16, 'out'))
    r16 = next(r for r in o16['results'] if r['band'][0] == '11th-20th')
    q16 = next(q for q in r16['per'])
    a16 = QC.check(os.path.join(tmp16, 'out'), r16['label'], q16['story'])
    b16 = QC.check(os.path.join(tmp16, 'out'), r16['label'], q16['story'], b_mm=500, h_mm=2000)
    check('T16 quick_check of the unchanged column = run (status, D/C); 500×2000 preview lowers D/C and flags the ETABS forces',
          a16['status'] == q16['status'] and abs(a16['dc'] - q16['dc']) < 0.01 and b16['dc'] < a16['dc'] and 'ETABS' in b16['note'],
          f"{a16['dc']} vs {q16['dc']:.3f}; 500x2000 {b16['dc']}")
finally:
    shutil.rmtree(tmp16, ignore_errors=True)

# ---------------------------------------------------------------- T17 f'c case table v1.5 (CEDA 27 Oct 2026)
o17 = FS.cases(M, cols15, sts, 400, [500, 600], dict(limits=[0.10, 0.15, 0.20, 0.25], fc_max=700))
c17 = o17['cases']
T17 = c17[0]
c17 = c17[1:]                                          # v1.5.1: row 0 = typical f'c in every story ('T')
check("T17 case table: T row first (400 everywhere, 2 failing = within 20 %), 500/1F (A=B, 1 failing, limit 10 %), stops at 600/1F "
      "with 0 failing; recommended at 10 % = 500/1F",
      T17['type'] == 'T' and T17['reduce_at'] is None and T17['lim'] == 0.20 and T17['n_fail'] == 2 and
      [(c['fc_bottom'], c['reduce_at'], c['type'], c['n_fail']) for c in c17] == [(500, '1F', 'A=B', 1), (600, '1F', 'A=B', 0)]
      and c17[0]['lim'] == 0.10 and o17['rec'] == {'10': 1, '15': 1, '20': 0, '25': 0} and o17['recommended'] == 1,
      str([(T17['fc_bottom'], T17['n_fail'], T17['lim'])] + [(c['fc_bottom'], c['reduce_at'], c['type'], c['n_fail'], c['lim']) for c in c17]))
o17b = FS.cases(M, cols15, sts, 400, [500, 600], dict(limits=[0.10, 0.15, 0.20, 0.25, 1.0], fc_max=700))
o17c = FS.cases(M, cols15, sts, 400, [500, 600], dict(limits=[0.10], fc_max=700))
check("T17 T row is always the first row, also when it fails every limit (limit 10 % only: T lim = None)",
      o17c['cases'][0]['type'] == 'T' and o17c['cases'][0]['lim'] is None and o17c['cases'][0]['n_fail'] == 2, str(o17c['cases'][0]['lim']))
check("T17 T row is recommended at every limit it satisfies (20 %, 25 %, 100 %) – f'c is not raised for nothing",
      o17b['rec'].get('20') == 0 and o17b['rec'].get('100') == 0 and o17b['cases'][0]['type'] == 'T', str(o17b['rec']))
check("T17 f'c-overstrength 9 · 90 % of the high-f'c concrete; low D/C 0.50–0.75 = 9 of 14 checked (typical + ρ ≤ 1.5 % not checked)",
      c17[0]['fcov_n'] == 9 and abs(c17[0]['fcov_pct'] - 90) < 1e-6 and c17[0]['checked'] == 14 and abs(c17[0]['low75_pct'] - 900 / 14) < 1e-6
      and c17[0]['low50_pct'] == 0, f"{c17[0]['fcov_n']} {c17[0]['fcov_pct']:.1f} {c17[0]['checked']} {c17[0]['low75_pct']:.1f}")

# ---------------------------------------------------------------- T18 input workbook v1.5 (ρ limits, SMF, enlarge / splice / marks)
tp = dict(criteria={'detailing.rho_ask_pct': 9, 'fc.limits_pct': '10, 20'}, stories=[], columns=[], sections=[], combos_rows=[], combo_cases=[],
          lap_rows=[], exclude=[], approve=[['A1', '1F']], overload=[], delta_ok=[], ranges=[], prices_rows=[],
          enlarge=[['A1', '1F', '2F', 'rect', 500, 900, 2]], splice=[['A1', '1F', '1F', 'Mechanical coupler']], marks=[['A1', 'C7']])
c18 = INPW.to_cfg(tp)
check('T18 ρ with approval clamped to 8 % (ACI 318-19 §10.6.1.1); Approve_rho used without size_restricted; limits read',
      c18['detailing']['rho_ask'] == 0.08 and c18['approved_rho_over_6'] == [['A1', '1F']] and c18['fc_rule15']['limits'] == [0.10, 0.20])
tp['criteria']['seismic.frame'] = 'SMF'
c18s = INPW.to_cfg(tp)
check('T18 SMF: ρ max and ρ with approval locked at 6 %, approvals ignored (ACI 318-19 §18.7.4.1)',
      c18s['detailing']['rho_max'] == 0.06 and c18s['detailing']['rho_ask'] == 0.06 and c18s['approved_rho_over_6'] == [])
tmp18 = tempfile.mkdtemp()
p18 = os.path.join(tmp18, 'w.xlsx'); open(p18, 'wb').write(INPW.make_template(None, tp))
r18 = INPW.read_template(p18)
c18r = INPW.to_cfg(r18)
check('T18 Enlarge / Splice / Marks sheets round trip -> section_overrides, splice_by, mark_names; no missing v1.5 keys',
      c18r['section_overrides'][0]['b_mm'] == 500 and c18r['section_overrides'][0]['option'] == 2 and c18r['splice_by'][0]['method'] == 'Mechanical coupler'
      and c18r['mark_names'] == {'A1': 'C7'} and r18['missing_keys'] == [], str(r18['missing_keys']))
shutil.rmtree(tmp18, ignore_errors=True)

# ---------------------------------------------------------------- T19 enlarged section, D/C target, splice, mark lock, proposals
import fix_tools as FX    # noqa: E402
tmp19 = tempfile.mkdtemp()
try:
    for f in ('c1_forces.json', 'c1_config.json'):
        shutil.copy(os.path.join(FIX, f), tmp19)
    cf = json.load(open(os.path.join(tmp19, 'c1_config.json')))
    cf['section_overrides'] = [dict(label='C1', story_from='B6-B5', story_to='1st-3th', shape='rect', b_mm=500, h_mm=2000, option=2)]
    cf['splice_by'] = [dict(label='C1', story_from='11th-20th', story_to='11th-20th', method='Mechanical coupler')]
    cf['mark_names'] = {'C1': 'C7'}
    cf['fix'] = dict(target_dc=0.90)
    json.dump(cf, open(os.path.join(tmp19, 'c19.json'), 'w'))
    o19 = run(os.path.join(tmp19, 'c19.json'), os.path.join(tmp19, 'out'))
    en = [r for r in o19['results'] if r.get('enlarged')]
    check('T19 enlarged rows: 500×2000 B6-B5…1st-3th, every story D/C ≤ 0.90 (target), ENLARGED / PROVISIONAL noted',
          len(en) == 4 and all(abs(r['b'] - 50) < 1e-9 for r in en) and all(q['dc'] <= 0.90 + 1e-6 for r in en for q in r['per'])
          and any(n.startswith('ENLARGED C1') and 'PROVISIONAL' in n for n in o19['notes']), str([(r['band'], [round(q['dc'], 3) for q in r['per']]) for r in en]))
    top = next(r for r in o19['results'] if r['band'][0] == '11th-20th')
    check('T19 splice per column: coupler -> no lap length; mark name locked (C7)', top['lap'] is None and top['splice'] == 'Mechanical coupler'
          and all(r['mark'] == 'C7' for r in o19['results']))
    p19 = FX.proposals(os.path.join(tmp19, 'out'), 'C1', ['4th-10th'], option=1, target=0.90)
    check('T19 fix_tools proposals: option 1 = the row of the failing story; every proposal D/C ≤ 0.90',
          p19['range'] == ['4th-10th'] and p19['options'] and all(o['dc_max'] <= 0.90 + 1e-6 for o in p19['options']),
          str([(o['kind'], o['b'], o['h'], o['dc_max']) for o in p19['options']]))
finally:
    shutil.rmtree(tmp19, ignore_errors=True)

# ---------------------------------------------------------------- T20 v1.5.1: steel members excluded, story order by top elevation
import etabs_io as E20
secs20 = {"C4 400x400  ( fc' 320 ksc )": dict(shape='rect', t3=40, t2=40)}
ex = [E20.non_rc(n, secs20) is not None for n in ('Truss Brace - H 150x150x7x10', 'Truss - H 350x350x12x19 (Top & Bottom)', 'W 12x26',
                                                   "C4 400x400  ( fc' 320 ksc )", 'C900x900fc400', 'CC 800', 'C 2000x450')]
check('T20 non-RC: steel truss / brace / W names excluded, concrete names kept (circle without its table still falls back to the name)',
      ex == [True, True, True, False, False, False, False], str(ex))
geo20 = {('a', '2 FL'): dict(zb=65, zt=695), ('b', '2 FL'): dict(zb=350, zt=695), ('c', 'Mezz'): dict(zb=65, zt=350),
         ('d', 'MDB'): dict(zb=65, zt=280), ('e', '1 FL'): dict(zb=-85, zt=65), ('f', 'Roof'): dict(zb=695, zt=1115)}
o20 = E20.story_order(geo20)
check('T20 story order by the story top, then the bottom: 1 FL, MDB, Mezz, 2 FL, Roof (partial-height stories from +0.65)',
      o20 == ['1 FL', 'MDB', 'Mezz', '2 FL', 'Roof'], str(o20))

print('ALL PASS' if not fails else f'{len(fails)} FAILED: {fails}')
sys.exit(1 if fails else 0)
