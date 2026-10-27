"""CEDA column input workbook (Excel) – every input the skill needs can come from this file.

Sheets (all optional except what the ETABS exports do not already give):
  Criteria   key | value | unit | note          design criteria (skill Step 1)
  Stories    story | z_bottom_m | z_top_m | basement | Lu_m | k | fc_ksc     bottom -> top
  Columns    label | story | x_m | y_m | Lu_m | k                         connectivity per column-story (optional)
  Sections   label | story_from | story_to | shape | b_mm | h_mm          b = t2 (My), h = t3 (Mx); blank stories = all
  Combos     combo | <load case> | <load case> | ...                      factors (header = ETABS load case names)
  Lap        fc_ksc | DB12 | DB16 | DB20 | DB25 | DB28 | DB32               Class B lap, mm (S-00-04-04)
  Exclude    combo                                                       combos not used
  Approve_rho label | story                                              ρ 6 % – approval limit approved (fix board, v1.5.0: no size_restricted)
  Enlarge    label | story_from | story_to | shape | b_mm | h_mm | option   enlarged sections from the fix board (v1.5.0)
  Splice     label | story_from | story_to | method                     splice method per column (default on Criteria)
  Marks      label | mark                                              mark names locked after the f'c case (v1.5.0)
  Accept_overload label | story | dc_max                                 D/C > 1 accepted by the engineer (Engineering judgment)
  Accept_delta label | story | delta_max                                 δ > 1.4 accepted by the engineer (Engineering judgment)
  Story_ranges story_from | story_to                                     story ranges fixed by the designer (v1.3.6; must cover
                                                                         every story once, bottom -> top; blank = engine ranges)
A CEDA development-length workbook (TABLE 11.04, one sheet per cover) can fill the Lap sheet: read_lap_table() (v1.3.6).
Template data are turned into ETABS-style tables (connectivity, section assignments, combo definitions) so the engine
reads them exactly like an ETABS export; where both exist, the input workbook wins (reported).
"""
import io, math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

CRITERIA = [  # key, default, unit, note
    ('project.name', '', '', 'Project name (header)'),
    ('project.location', '', '', ''),
    ('project.building', '', '', ''),
    ('project.job', 'JOB', '', 'Job no. – also the file prefix'),
    ('project.date', '', '', ''),
    ('project.calc_by_1', '', '', 'name, licence'),
    ('project.calc_by_2', '', '', 'name, licence'),
    ('project.approve_by', '', '', 'name, licence'),
    ('seismic.design', 'Y', 'Y/N', 'Seismic design required'),
    ('seismic.zone', '', '', ''),
    ('seismic.importance', '', '', 'Importance factor I'),
    ('seismic.category', 'ข', 'ก/ข/ค/ง', 'Seismic design category'),
    ('seismic.system', 'Building frame + ordinary RC shear wall', '', 'Lateral system (มยผ. 1301/1302-61 Table 2.3-1)'),
    ('seismic.R', 5, '', ''),
    ('seismic.Omega0', 2.5, '', ''),
    ('seismic.Cd', 4.5, '', ''),
    ('seismic.frame', 'gravity', 'gravity/IMF/SMF', 'Columns are'),
    ('materials.fc_typical', 400, 'ksc', "Typical f'c of the project (cylinder) – columns reduce to it"),
    ('materials.fy', 5000, 'ksc', ''),
    ('materials.fyt', 4000, 'ksc', ''),
    ('materials.fc_slab', 320, 'ksc', 'ACI 15.5 check'),
    ('materials.fc_mode', 'auto', 'auto/fixed', "auto = skill f'c study; fixed = fc_bottom up to fc_reduce_at (or Stories fc_ksc)"),
    ('materials.fc_bottom', 560, 'ksc', 'fixed mode only'),
    ('materials.fc_reduce_at', '', 'story', 'fixed mode only – first story with the typical f\'c'),
    ('detailing.cover_mm', 40, 'mm', 'clear cover to ties'),
    ('detailing.tie_db', 12, 'mm', '12 or 16'),
    ('detailing.rho_max_pct', 6.0, '%', 'ρ max without approval – every column, steps 3–4 (CEDA practice)'),
    ('detailing.rho_ask_pct', 7.0, '%', 'ρ max with approval on the fix board – adjustable 6–8 (ACI 318-19 §10.6.1.1 ≤ 8 %). SMF: both locked at 6 % (ACI 318-19 §18.7.4.1)'),
    ('detailing.splice', 'Lap – staggered', 'Lap – staggered / Lap – one level / Mechanical coupler', 'project default; per column on the Splice sheet'),
    ('fix.target_dc', 0.90, '', 'enlarged sections: proposals and typed sizes must reach D/C ≤ this with the old forces (ASSUMPTION, 10 % margin)'),
    ('fix.step_mm', 50, 'mm', 'enlargement step'),
    ('fix.max_add_mm', 500, 'mm', 'largest increase per side (ASSUMPTION)'),
    ('fix.compare_ddc', 0.05, '', 'round compare: show |ΔD/C| above this (ASSUMPTION)'),
    ('marks.renumber', 'Y', 'Y/N', 'at FINAL: rename marks by section size (names only)'),
    ('slender.k', 1.0, '', 'default k (Stories / Columns may override)'),
    ('slender.beta_d', '', '', 'blank = from DL/SDL/LL load cases (skill Rev.1)'),
    ('slender.sway', '', '', 'blank = check Q from Story Forces/Drifts; non-sway = assumed by the engineer (Judgment)'),
    ('gravity.DL', 'DL, Dead, D, SW', '', 'ETABS load case names counted as DL'),
    ('gravity.SDL', 'SDL, SIDL, SuperDead', '', 'ETABS load case names counted as SDL'),
    ('gravity.LL', 'LL, Live, L, Lr', '', 'ETABS load case names counted as LL'),
    ('marks.envelope', 'Y', 'Y/N', ''),
    ('marks.max_rows', None, '', 'rows that may differ (blank = no limit, CEDA criterion A)'),
    ('marks.max_step', None, '', 'ladder steps per row (blank = no limit)'),
    ('marks.project_max_weight_pct', 10, '%', 'extra steel of the whole project from mark envelopes'),
    ('marks.max_weight_pct', 25, '%', 'extra longitudinal steel of each label'),
    ('story_ranges.mode', 'auto', 'auto/each', 'auto = basement every story, above ground 2–3 stories / 12 m + Lst'),
    ('fc.limits_pct', '10, 15, 20, 25', '%', "f'c study v1.5: failing limits per story, count rule max(limit × columns, 1)"),
    ('fc.max', 700, 'ksc', "highest f'c of the study (development-table values only; > 700 rarely used in Thailand)"),
    ('fc.dc_min', 0.75, '', "under-use threshold (information only)"),
    ('fc.dc_low', 0.50, '', "low-D/C band (information only)"),
    ('fc.exempt_rho_pct', 1.5, '%', "typical-f'c column-stories with ρ ≤ this are not checked for low D/C"),
    ('fc.share_low_pct', 50, '%', "reduction story: share of columns at ρ ≤ 1.5 % with the typical f'c"),
]
HDR = {
    'Stories': ['story', 'z_bottom_m', 'z_top_m', 'basement', 'Lu_m', 'k', 'fc_ksc'],
    'Columns': ['label', 'story', 'x_m', 'y_m', 'Lu_m', 'k'],
    'Sections': ['label', 'story_from', 'story_to', 'shape', 'b_mm', 'h_mm'],
    'Lap': ['fc_ksc', 'DB12', 'DB16', 'DB20', 'DB25', 'DB28', 'DB32'],
    'Exclude': ['combo'],
    'Approve_rho': ['label', 'story'],
    'Enlarge': ['label', 'story_from', 'story_to', 'shape', 'b_mm', 'h_mm', 'option'],
    'Splice': ['label', 'story_from', 'story_to', 'method'],
    'Marks': ['label', 'mark'],
    'Accept_overload': ['label', 'story', 'dc_max'],
    'Accept_delta': ['label', 'story', 'delta_max'],
    'Story_ranges': ['story_from', 'story_to'],
    'Prices': ['item', 'size', 'unit', 'price_thb'],
}
NOTE = {
    'Stories': 'Bottom → top. z in m (gives story order and Lu = z_top − z_bottom). Lu_m / k override for the whole story. '
               'basement = Y for stories below the seismic model. fc_ksc: optional fixed f\'c per story.',
    'Columns': 'Optional – one row per column-story: XY (m) to match labels between models, Lu_m / k override for that column.',
    'Sections': 'b = t2 (bending My), h = t3 (bending Mx), mm. shape rect/circ (circ: b = h = D). Blank story_from/to = all stories.',
    'Combos': 'One row per combo; columns after "combo" = ETABS load case names (edit the headers), cells = factors.',
    'Lap': 'Class B tension lap (mm) from S-00-04-04 of the project. Blank = unknown (never invented).',
    'Exclude': 'Combos not used in the design (e.g. cross-standard U_B combos).',
    'Approve_rho': 'Label-stories for which ρ above 6 % (up to the approval limit on Criteria, ≤ 8 %) is approved on the fix board. Not used in the f\'c study. Ignored for SMF.',
    'Enlarge': 'Enlarged sections from the fix board: stories from–to (bottom → top), new b × h (mm; circ: b = D), option 1 = row only / 2 = failing story down to the lowest. Forces stay from the ETABS export until it is re-run (PROVISIONAL).',
    'Splice': 'Splice method per column range: Lap – staggered / Lap – one level / Mechanical coupler. Blank = Criteria default.',
    'Marks': 'Mark names locked after the f\'c case is chosen; new groups get new numbers. Filled by the app.',
    'Accept_overload': 'Label-stories whose D/C above 1.0 the engineer accepts (up to dc_max) – Engineering judgment, reported; ACI 318 needs φSn ≥ U.',
    'Accept_delta': 'Label-stories whose moment magnifier above 1.4 the engineer accepts (up to delta_max) – Engineering judgment; ACI 6.2.5.3 limits it to 1.4.',
    'Story_ranges': 'Optional – story ranges fixed by the designer, bottom → top, every story exactly once (e.g. +00 → GF merged). Blank = engine ranges.',
    'Prices': "Optional unit prices for the f'c what-if cost (material only): concrete per f'c (THB/m³), longitudinal steel per bar size (THB/kg).",
}
PRICE_ROWS = [['concrete', f, 'THB/m3', None] for f in (240, 280, 320, 400, 450, 500, 560, 600, 700)] + \
             [['steel', f'DB{d}', 'THB/kg', None] for d in (12, 16, 20, 25, 28, 32)]
BLUE = Font(name='Arial', size=10, color='0000FF')
B = Font(name='Arial', size=10, bold=True)
N = Font(name='Arial', size=10)
HFILL = PatternFill('solid', fgColor='D9E1F2')
LAP_KNOWN = {400: {25: 1550, 28: 1875, 32: 2325}, 560: {25: 1325, 28: 1600, 32: 1975}}


def _sheet(wb, name, hdr, rows, note):
    ws = wb.create_sheet(name)
    ws['A1'] = note; ws['A1'].font = Font(name='Arial', size=9, italic=True)
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=2, column=j, value=h); c.font = B; c.fill = HFILL
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = max(12, len(str(h)) + 4)
    for i, r in enumerate(rows, 3):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v); c.font = BLUE
    ws.freeze_panes = 'A3'
    return ws


def make_template(found=None, current=None):
    """found: dict(labels, stories, combos, cases, Lu, sec_of, z) from the ETABS uploads; current: read_template() dict to keep."""
    found = found or {}
    cur = current or {}
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'Criteria'
    ws['A1'] = 'CEDA column design – input workbook. Blue cells are inputs. Every sheet is optional where the ETABS export already gives the data.'
    ws['A1'].font = Font(name='Arial', size=9, italic=True)
    for j, h in enumerate(['key', 'value', 'unit', 'note'], 1):
        c = ws.cell(row=2, column=j, value=h); c.font = B; c.fill = HFILL
    cv = cur.get('criteria', {})
    for i, (k, d, u, n) in enumerate(CRITERIA, 3):
        ws.cell(row=i, column=1, value=k).font = N
        c = ws.cell(row=i, column=2, value=cv.get(k, d)); c.font = BLUE
        ws.cell(row=i, column=3, value=u).font = N
        ws.cell(row=i, column=4, value=n).font = N
    for col, w in zip('ABCD', (26, 40, 14, 80)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A3'
    # Stories
    st_rows = cur.get('stories') or [[s, *(found.get('z', {}).get(s, (None, None))), None, None, None, None] for s in found.get('stories', [])]
    _sheet(wb, 'Stories', HDR['Stories'], st_rows, NOTE['Stories'])
    _sheet(wb, 'Columns', HDR['Columns'], cur.get('columns') or [], NOTE['Columns'])
    sec_rows = cur.get('sections')
    if not sec_rows:
        sec_rows = []
        for lab in found.get('labels', []):
            s = found.get('sec_of', {}).get(lab)
            if s:
                for shape, b, h, _ in sorted(s):
                    sec_rows.append([lab, None, None, shape, b, h])
            else:
                sec_rows.append([lab, None, None, 'rect', None, None])
    ws2 = _sheet(wb, 'Sections', HDR['Sections'], sec_rows, NOTE['Sections'])
    dv = DataValidation(type='list', formula1='"rect,circ"', allow_blank=True); ws2.add_data_validation(dv); dv.add('D3:D2000')
    cases = cur.get('combo_cases') or found.get('cases') or ['DL', 'SDL', 'LL', 'EQX', 'EQY', 'WX', 'WY']
    cmb_rows = cur.get('combos_rows') or [[c] + [None] * len(cases) for c in found.get('undefined', found.get('combos', []))]
    _sheet(wb, 'Combos', ['combo'] + list(cases), cmb_rows, NOTE['Combos'])
    lap_rows = cur.get('lap_rows') or [[f] + [LAP_KNOWN.get(f, {}).get(d) for d in (12, 16, 20, 25, 28, 32)] for f in (400, 560)]
    _sheet(wb, 'Lap', HDR['Lap'], lap_rows, NOTE['Lap'])
    _sheet(wb, 'Exclude', HDR['Exclude'], [[c] for c in cur.get('exclude', [])], NOTE['Exclude'])
    _sheet(wb, 'Approve_rho', HDR['Approve_rho'], cur.get('approve', []), NOTE['Approve_rho'])
    _sheet(wb, 'Enlarge', HDR['Enlarge'], cur.get('enlarge', []), NOTE['Enlarge'])
    _sheet(wb, 'Splice', HDR['Splice'], cur.get('splice', []), NOTE['Splice'])
    _sheet(wb, 'Marks', HDR['Marks'], cur.get('marks', []), NOTE['Marks'])
    _sheet(wb, 'Accept_overload', HDR['Accept_overload'], cur.get('overload', []), NOTE['Accept_overload'])
    _sheet(wb, 'Accept_delta', HDR['Accept_delta'], cur.get('delta_ok', []), NOTE['Accept_delta'])
    _sheet(wb, 'Story_ranges', HDR['Story_ranges'], cur.get('ranges', []), NOTE['Story_ranges'])
    _sheet(wb, 'Prices', HDR['Prices'], cur.get('prices_rows') or PRICE_ROWS, NOTE['Prices'])
    bio = io.BytesIO(); wb.save(bio)
    return bio.getvalue()


def _blank(v):
    return v is None or (isinstance(v, str) and not v.strip()) or (isinstance(v, float) and math.isnan(v))


def _rows(ws, ncol=None):
    hdr = [c.value for c in ws[2]]
    hdr = [str(h).strip() if h is not None else '' for h in hdr]
    while hdr and not hdr[-1]:
        hdr.pop()
    out = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        r = list(r)[:len(hdr)]
        if all(_blank(v) for v in r):
            continue
        out.append(r + [None] * (len(hdr) - len(r)))
    return hdr, out


def is_template(path):
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        return 'Criteria' in wb.sheetnames or ('Sections' in wb.sheetnames and 'Stories' in wb.sheetnames)
    except Exception:
        return False


def read_template(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = dict(criteria={}, stories=[], columns=[], sections=[], combos_rows=[], combo_cases=[], lap_rows=[], exclude=[], approve=[], overload=[], delta_ok=[], ranges=[], prices_rows=[], enlarge=[], splice=[], marks=[], issues=[], missing_keys=[])
    if 'Criteria' in wb.sheetnames:
        for r in wb['Criteria'].iter_rows(min_row=3, values_only=True):
            if r and r[0]:
                out['criteria'][str(r[0]).strip()] = r[1]
    for name, key in (('Stories', 'stories'), ('Columns', 'columns'), ('Sections', 'sections'), ('Lap', 'lap_rows'), ('Approve_rho', 'approve'), ('Accept_overload', 'overload'), ('Accept_delta', 'delta_ok'), ('Story_ranges', 'ranges'), ('Prices', 'prices_rows'), ('Enlarge', 'enlarge'), ('Splice', 'splice'), ('Marks', 'marks')):
        if name in wb.sheetnames:
            hdr, rows = _rows(wb[name])
            exp = HDR[name]
            if hdr[:len(exp)] != exp:
                out['issues'].append(f'{name}: header {hdr} differs from {exp} – columns read by position')
            out[key] = [r[:len(exp)] + [None] * (len(exp) - len(r)) for r in rows]
    if 'Combos' in wb.sheetnames:
        hdr, rows = _rows(wb['Combos'])
        out['combo_cases'] = hdr[1:]
        out['combos_rows'] = rows
    if 'Exclude' in wb.sheetnames:
        out['exclude'] = [str(r[0]).strip() for r in _rows(wb['Exclude'])[1] if not _blank(r[0])]
    out['sections'] = [r for r in out['sections'] if not _blank(r[0])]
    out['missing_keys'] = [k for k, *_ in CRITERIA if k not in out['criteria']]   # v1.5.0: defaults used (reported on Check)
    return out


def _num(v):
    try:
        return None if _blank(v) else float(v)
    except Exception:
        return None


def story_list(tpl):
    return [str(r[0]).strip() for r in tpl.get('stories', []) if not _blank(r[0])]


def to_tables(tpl, label_story):
    """template -> ETABS-style tables {name: rows}, units, issues.  label_story: set of (label, story) in the forces."""
    T, U, issues = {}, {}, []
    st = {str(r[0]).strip(): r for r in tpl.get('stories', []) if not _blank(r[0])}
    order = story_list(tpl)
    cols = {(str(r[0]).strip(), str(r[1]).strip()): r for r in tpl.get('columns', []) if not _blank(r[0]) and not _blank(r[1])}
    keys = sorted(set(label_story) | set(cols))
    # connectivity (always when Stories or Columns are given)
    if st or cols or tpl.get('sections'):
        pts, cn = [], []
        pid = 1
        for lab, s in keys:
            c = cols.get((lab, s))
            srow = st.get(s)
            zb, zt = (_num(srow[1]), _num(srow[2])) if srow else (None, None)
            Lu = _num(c[4]) if c else None
            L = Lu if Lu else ((zt - zb) if zb is not None and zt is not None else (_num(srow[4]) if srow else None))
            row = {'Unique Name': f'T{lab}_{s}', 'Story': s, 'ColumnBay': lab}
            x, y = (_num(c[2]), _num(c[3])) if c else (None, None)
            if zb is not None and zt is not None and x is not None and y is not None:
                # points only with real XY: without XY the engine matches models by label and orders stories by Stories
                pts.append({'Unique Name': f'P{pid}', 'Story': s, 'X': x, 'Y': y, 'Z': zb})
                pts.append({'Unique Name': f'P{pid + 1}', 'Story': s, 'X': x, 'Y': y, 'Z': zt})
                row.update({'UniquePtI': f'P{pid}', 'UniquePtJ': f'P{pid + 1}'}); pid += 2
            if L:
                row['Length'] = L
            cn.append(row)
        T['Column Object Connectivity'] = cn; U['Column Object Connectivity'] = {'Length': 'm'}
        if pts:
            T['Point Object Connectivity'] = pts; U['Point Object Connectivity'] = {'X': 'm', 'Y': 'm', 'Z': 'm'}
    # sections
    if tpl.get('sections'):
        idx = {s: i for i, s in enumerate(order)}
        allst = sorted({s for _, s in keys}, key=lambda s: idx.get(s, 1e9))
        defs, asg = {}, []
        for r in tpl['sections']:
            lab = str(r[0]).strip(); s0 = None if _blank(r[1]) else str(r[1]).strip(); s1 = None if _blank(r[2]) else str(r[2]).strip()
            shape = (str(r[3]).strip().lower() if not _blank(r[3]) else 'rect')
            b, h = _num(r[4]), _num(r[5])
            if b is None or (shape == 'rect' and h is None):
                issues.append(f'Sections: {lab} has no b/h'); continue
            if shape == 'circ':
                h = b
            name = f'T_{shape}_{b:.0f}x{h:.0f}'
            defs[name] = (shape, b, h)
            for s in allst:
                if (lab, s) not in set(keys):
                    continue
                if s0 and idx.get(s, -1) < idx.get(s0, -1):
                    continue
                if s1 and idx.get(s, 1e9) > idx.get(s1, 1e9):
                    continue
                asg.append({'Story': s, 'Label': lab, 'UniqueName': f'T{lab}_{s}', 'Section Property': name})
        T['Frame Assignments - Section Properties'] = asg; U['Frame Assignments - Section Properties'] = {}
        rect = [{'Name': n, 'Depth': h / 1000, 'Width': b / 1000} for n, (sh, b, h) in defs.items() if sh == 'rect']
        circ = [{'Name': n, 'Diameter': b / 1000} for n, (sh, b, h) in defs.items() if sh == 'circ']
        if rect:
            T['Frame Section Property Definitions - Concrete Rectangular'] = rect
            U['Frame Section Property Definitions - Concrete Rectangular'] = {'Depth': 'm', 'Width': 'm'}
        if circ:
            T['Frame Section Property Definitions - Concrete Circle'] = circ
            U['Frame Section Property Definitions - Concrete Circle'] = {'Diameter': 'm'}
    # combos
    if tpl.get('combos_rows'):
        rows = []
        for r in tpl['combos_rows']:
            if _blank(r[0]):
                continue
            for case, v in zip(tpl['combo_cases'], r[1:]):
                f = _num(v)
                if f:
                    rows.append({'Name': str(r[0]).strip(), 'Load Name': str(case).strip(), 'SF': f})
        if rows:
            T['Load Combination Definitions'] = rows; U['Load Combination Definitions'] = {}
    return T, U, issues


def AIO_LEN(u):
    """metres per unit of the length column of an ETABS section table (template rows are in m)."""
    import etabs_io as E
    unit = next((v for k, v in u.items() if k in ('Depth', 'Width', 'Diameter', 't3', 't2') and v), 'm')
    return E.unit_factor(unit, 'length') / 100.0


def apply(tables, units, tpl, label_story):
    """merge the template into one model's ETABS tables. Explicit template data win; returns new tables, units, report."""
    TT, UU, issues = to_tables(tpl, label_story)
    T, U = dict(tables), dict(units)
    rep = list(issues)
    has = lambda n: bool(T.get(n))
    # connectivity: export without connectivity -> template (Stories / Columns / Sections) gives it;
    # export with connectivity -> only the column-stories listed in Columns are replaced (others keep ETABS XY / length)
    if 'Column Object Connectivity' in TT:
        if not has('Column Object Connectivity'):
            for n in ('Column Object Connectivity', 'Point Object Connectivity'):
                if n in TT:
                    T[n], U[n] = TT[n], UU[n]
        elif tpl.get('columns'):
            mine = {(str(r[0]).strip(), str(r[1]).strip()) for r in tpl['columns'] if not _blank(r[0]) and not _blank(r[1])}
            import etabs_io as E
            keep = [r for r in T['Column Object Connectivity']
                    if (str(E.col(r, 'Column Bay', 'ColumnBay', 'Label', 'Column')), str(E.col(r, 'Story'))) not in mine]
            fL = E.unit_factor((U.get('Column Object Connectivity') or {}).get('Length', 'm'), 'length') / 100.0   # m per unit
            new_rows = [dict(r, Length=r['Length'] / fL) if 'Length' in r else r
                        for r in TT['Column Object Connectivity'] if (r['ColumnBay'], r['Story']) in mine]
            T['Column Object Connectivity'] = keep + new_rows
            if 'Point Object Connectivity' in TT:
                fP = E.unit_factor((U.get('Point Object Connectivity') or {}).get('X', 'm'), 'length') / 100.0 if T.get('Point Object Connectivity') else 1.0
                T['Point Object Connectivity'] = list(T.get('Point Object Connectivity', [])) + [
                    {k: (v / fP if k in ('X', 'Y', 'Z') else v) for k, v in r.items()} for r in TT['Point Object Connectivity']]
                U.setdefault('Point Object Connectivity', UU['Point Object Connectivity'])
            rep.append(f'Connectivity: input workbook Columns replace {len(new_rows)} ETABS column-stories')
    # sections: per label – labels listed in the workbook replace the ETABS assignment, other labels keep ETABS
    if 'Frame Assignments - Section Properties' in TT:
        # v1.3.3: per label-story (a Sections row with story_from / story_to replaces only those stories)
        mine_ls = {(str(r['Label']), str(r['Story'])) for r in TT['Frame Assignments - Section Properties']}
        mine = {l for l, _ in mine_ls}
        old = T.get('Frame Assignments - Section Properties', [])
        keep = [r for r in old if (str(r.get('Label') or r.get('Column Bay') or ''), str(r.get('Story') or '')) not in mine_ls]
        if len(keep) != len(old):
            rep.append(f'Column sizes: input workbook Sections replace {len(old) - len(keep)} ETABS label-stories of ' +
                       ', '.join(sorted(mine)[:10]) + (' …' if len(mine) > 10 else ''))
        T['Frame Assignments - Section Properties'] = keep + TT['Frame Assignments - Section Properties']
        U.setdefault('Frame Assignments - Section Properties', {})
        for n in [k for k in TT if k.startswith('Frame Section Property Definitions')]:
            ex = T.get(n, [])
            T[n] = ex + [r for r in TT[n] if r['Name'] not in {e.get('Name') for e in ex}]
            U[n] = U.get(n) or UU[n]
            if ex and U[n] != UU[n]:          # units of the ETABS table differ: convert the template rows
                fx = 1.0 / AIO_LEN(U[n])
                T[n] = ex + [{k: (v * fx if k != 'Name' else v) for k, v in r.items()} for r in TT[n] if r['Name'] not in {e.get('Name') for e in ex}]
    # combos: template rows win per combo
    if 'Load Combination Definitions' in TT:
        mine = {r['Name'] for r in TT['Load Combination Definitions']}
        old = [r for r in T.get('Load Combination Definitions', []) if str(r.get('Name') or '') not in mine]
        if len(old) != len(T.get('Load Combination Definitions', [])):
            rep.append('Load combinations: input workbook factors replace the ETABS definition for ' + ', '.join(sorted(mine)[:6]))
        # ETABS combo tables use blank names on continuation rows – fill them before filtering
        cur, filled = None, []
        for r in T.get('Load Combination Definitions', []):
            nm = r.get('Name') or cur; cur = nm
            if str(nm) not in mine:
                filled.append(dict(r, Name=nm))
        T['Load Combination Definitions'] = filled + TT['Load Combination Definitions']
        U['Load Combination Definitions'] = U.get('Load Combination Definitions', {})
    return T, U, rep


def to_cfg(tpl):
    """criteria + per-story / per-column overrides -> partial engine config."""
    c = {k: v for k, v in tpl.get('criteria', {}).items()}
    g = lambda k, d=None: d if _blank(c.get(k)) else c.get(k)
    yn = lambda k, d='N': str(g(k, d)).strip().upper().startswith('Y')
    names = lambda k: [x.strip() for x in str(g(k, '')).split(',') if x.strip()]
    split = lambda s: [x.strip() for x in (str(s or '').split(',', 1) + [''])[:2]]
    P = dict(name=g('project.name', ''), location=g('project.location', ''), building=g('project.building', ''),
             job=str(g('project.job', 'JOB')), date=str(g('project.date', '')),
             calc_by=[split(g('project.calc_by_1')), split(g('project.calc_by_2'))], approve_by=split(g('project.approve_by')))
    seis = dict(design=yn('seismic.design', 'Y'), zone=g('seismic.zone', ''), importance=g('seismic.importance', ''),
                category=str(g('seismic.category', 'ข')), system=g('seismic.system', ''), R=_num(g('seismic.R')),
                Omega0=_num(g('seismic.Omega0')), Cd=_num(g('seismic.Cd')), frame=str(g('seismic.frame', 'gravity')), imf_E_factor=2.0)
    order = story_list(tpl)
    mat = dict(fy=_num(g('materials.fy', 5000)), fyt=_num(g('materials.fyt', 4000)), fc_typical=_num(g('materials.fc_typical', 400)),
               fc_slab=_num(g('materials.fc_slab', 320)), fc_list=[210, 240, 280, 320, 400, 450, 500, 560, 600, 700])
    fcs = {str(r[0]).strip(): _num(r[6]) for r in tpl.get('stories', []) if not _blank(r[0])}
    if str(g('materials.fc_mode', 'auto')).lower().startswith('fix'):
        if fcs and all(v for v in fcs.values()):
            mat['fc_by_story'] = fcs
        elif order and g('materials.fc_reduce_at'):
            ra = order.index(str(g('materials.fc_reduce_at')).strip())
            mat['fc_by_story'] = {s: (_num(g('materials.fc_bottom', 560)) if i < ra else mat['fc_typical']) for i, s in enumerate(order)}
    ov = []
    for r in tpl.get('stories', []):
        if _blank(r[0]):
            continue
        Lu, k = _num(r[4]), _num(r[5])
        if Lu or k:
            o = dict(stories=[str(r[0]).strip()], labels=[])
            if Lu: o['Lu_m'] = Lu
            if k: o['k'] = k
            ov.append(o)
    for r in tpl.get('columns', []):
        if _blank(r[0]) or _blank(r[1]):
            continue
        Lu, k = _num(r[4]), _num(r[5])
        if Lu or k:
            o = dict(stories=[str(r[1]).strip()], labels=[str(r[0]).strip()])
            if Lu: o['Lu_m'] = Lu
            if k: o['k'] = k
            ov.append(o)
    lap = {}
    for r in tpl.get('lap_rows', []):
        f = _num(r[0])
        if f:
            lap[str(int(f))] = {d: _num(v) for d, v in zip(('12', '16', '20', '25', '28', '32'), r[1:]) if _num(v)}
    smf = 'SMF' in str(g('seismic.frame', '')).upper() or 'SPECIAL MOMENT' in str(g('seismic.system', '')).upper()
    rmax = min((_num(g('detailing.rho_max_pct', 6.0)) or 6.0) / 100, 0.06)
    rask = min(max((_num(g('detailing.rho_ask_pct', 7.0)) or 7.0) / 100, rmax), 0.08)        # ACI 318-19 §10.6.1.1
    if smf:                                                                                   # ACI 318-19 §18.7.4.1
        rmax = rask = 0.06
    splice_def = str(g('detailing.splice', 'Lap – staggered'))
    lims = [float(x) / 100 for x in str(g('fc.limits_pct', '10, 15, 20, 25')).replace(';', ',').split(',') if x.strip()] or [0.10, 0.15, 0.20, 0.25]
    gn = {'DL': names('gravity.DL') or ['DL', 'Dead', 'D', 'SW'], 'SDL': names('gravity.SDL') or ['SDL', 'SIDL', 'SuperDead'],
          'LL': names('gravity.LL') or ['LL', 'Live', 'L', 'Lr']}
    cfg = dict(project=P, seismic=seis, materials=mat,
               detailing=dict(cover_mm=_num(g('detailing.cover_mm', 40)), tie_db=int(_num(g('detailing.tie_db', 12))), kmax=4,
                              cc=[125, 150], rho_max=rmax, rho_ask=rask, so_mm=150, s_mm=300, rho_min=0.01, splice=splice_def,
                              small_face=[[200, 2], [540, 3], [660, 4]]),
               slender=dict(k=_num(g('slender.k', 1.0)) or 1.0, beta_d=_num(g('slender.beta_d')), overrides=ov, Lu_default_m=3.0, sway=str(g('slender.sway', '') or '')),
               gravity_cases=gn, lap_table=lap, exclude_combos=list(tpl.get('exclude', [])),
               approved_rho_over_6=[] if smf else [[str(a), str(b)] for a, b in tpl.get('approve', []) if not _blank(a) and not _blank(b)],
               section_overrides=[dict(label=str(r[0]).strip(), story_from=str(r[1]).strip(), story_to=str(r[2]).strip(),
                                       shape=str(r[3] or 'rect').strip(), b_mm=_num(r[4]), h_mm=_num(r[5]) or _num(r[4]), option=int(_num(r[6]) or 2))
                                  for r in tpl.get('enlarge', []) if not _blank(r[0]) and _num(r[4])],
               splice_by=[dict(label=str(r[0]).strip(), story_from=str(r[1] or '').strip(), story_to=str(r[2] or '').strip(), method=str(r[3]).strip())
                          for r in tpl.get('splice', []) if not _blank(r[0]) and not _blank(r[3])],
               mark_names={str(r[0]).strip(): str(r[1]).strip() for r in tpl.get('marks', []) if not _blank(r[0]) and not _blank(r[1])},
               mark_renumber=yn('marks.renumber', 'Y'),
               fix=dict(target_dc=_num(g('fix.target_dc', 0.90)) or 0.90, step_mm=_num(g('fix.step_mm', 50)) or 50,
                        max_add_mm=_num(g('fix.max_add_mm', 500)) or 500, compare_ddc=_num(g('fix.compare_ddc', 0.05)) or 0.05),
               accept_overload=[[str(a), str(b), _num(c)] for a, b, c in tpl.get('overload', []) if not _blank(a) and not _blank(c)],
               accept_delta=[[str(a), str(b), _num(c)] for a, b, c in tpl.get('delta_ok', []) if not _blank(a) and not _blank(c)],
               story_ranges=dict(mode=str(g('story_ranges.mode', 'auto')), max_group=3, stock_m=12.0, lst_m=None,
                                 groups=[[str(a).strip(), str(b).strip()] for a, b in tpl.get('ranges', []) if not _blank(a) and not _blank(b)]),
               basement_stories=[str(r[0]).strip() for r in tpl.get('stories', []) if not _blank(r[0]) and str(r[3] or '').upper().startswith('Y')],
               story_order=order,
               fc_rule15=dict(limits=lims, fc_max=_num(g('fc.max', 700)) or 700, dc_min=_num(g('fc.dc_min', 0.75)) or 0.75,
                              dc_low=_num(g('fc.dc_low', 0.50)) or 0.50, exempt_rho=(_num(g('fc.exempt_rho_pct', 1.5)) or 1.5) / 100,
                              share_low=(_num(g('fc.share_low_pct', 50)) or 50) / 100, rho_low=0.015, rho_max=rmax, rho_ask=rask),
               prices=prices_dict(tpl.get('prices_rows')),
               mark_envelope=dict(envelope=yn('marks.envelope', 'Y'),
                                  max_rows=None if _blank(g('marks.max_rows', None)) else int(_num(g('marks.max_rows', None))),
                                  max_step=None if _blank(g('marks.max_step', None)) else int(_num(g('marks.max_step', None))),
                                  max_weight=_num(g('marks.max_weight_pct', 25)) / 100,
                                  project_max_weight=_num(g('marks.project_max_weight_pct', 10)) / 100))
    return cfg


def check_ranges(groups, order):
    """v1.3.6: designer story ranges must cover every story of the order exactly once, bottom -> top. -> list of problems."""
    if not groups:
        return []
    idx = {s: i for i, s in enumerate(order)}
    bad = [s for g in groups for s in g if s not in idx]
    if bad:
        return [f'Story_ranges: story not in the story order: {", ".join(sorted(set(bad)))}']
    gs = sorted(groups, key=lambda g: idx[g[0]])
    out, nxt = [], 0
    for a, b in gs:
        if idx[b] < idx[a]:
            out.append(f'Story_ranges: {a} → {b} runs top → bottom')
        if idx[a] != nxt:
            out.append(f'Story_ranges: gap or overlap before {a} (expected {order[nxt] if nxt < len(order) else "end"})')
        nxt = max(nxt, idx[b] + 1)
    if nxt != len(order):
        out.append(f'Story_ranges: stories above {order[nxt - 1]} not covered')
    return out


def read_lap_table(path, sheet=None, col='Lst'):
    """v1.3.6: CEDA development-length workbook 'TABLE 11.04 : STANDARD DEVELOPMENT LENGTH' (one sheet per cover, e.g. '40mm')
    -> Lap rows [[fc, DB12, DB16, DB20, DB25, DB28, DB32], ...] of Class B tension lap 'Lst' (other bar), mm, plus a note.
    Columns are found by the header labels (L*dt, Ldt, Ldh, Ldc, L*st, Lst, Lsc); f'c is filled down; bar size from column A."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    hdr_row, ci = None, None
    for i, r in enumerate(ws.iter_rows(max_row=15, values_only=True), 1):
        labs = [str(v).strip() if v is not None else '' for v in r]
        if col in labs:
            hdr_row, ci = i, labs.index(col)
            break
    if hdr_row is None:
        raise ValueError(f'{ws.title}: header "{col}" not found – not a CEDA TABLE 11.04 sheet')
    tab, fc = {}, None
    for r in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
        r = list(r) + [None] * 12
        if isinstance(r[1], (int, float)):
            fc = int(r[1])
        d, v = r[0], r[ci]
        if fc and isinstance(d, (int, float)) and int(d) in (12, 16, 20, 25, 28, 32) and isinstance(v, (int, float)):
            tab.setdefault(fc, {})[int(d)] = float(v)
    rows = [[f] + [tab[f].get(d) for d in (12, 16, 20, 25, 28, 32)] for f in sorted(tab)]
    return rows, f'{len(rows)} f\'c rows from "{ws.title}" ({col}, Class B other bar)'


def lap_sheets(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    return list(wb.sheetnames)


def prices_dict(rows):
    """Prices rows -> {'concrete': {'400': THB/m3}, 'steel': {'32': THB/kg}} or None when no price is given (v1.4.0)."""
    out = {'concrete': {}, 'steel': {}}
    for r in rows or []:
        if not r or _blank(r[0]) or _blank(r[1]) or _num(r[3] if len(r) > 3 else None) is None:
            continue
        item = str(r[0]).strip().lower()
        key = str(r[1]).strip().upper().replace('DB', '')
        try:
            key = str(int(float(key)))
        except ValueError:
            continue
        if item.startswith('conc'):
            out['concrete'][key] = _num(r[3])
        elif item.startswith('steel') or item.startswith('rebar'):
            out['steel'][key] = _num(r[3])
    return out if (out['concrete'] or out['steel']) else None
