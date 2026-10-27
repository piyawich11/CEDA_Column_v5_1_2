"""Draft TABLE OF COLUMN – PDF + DXF (CEDA Revit-export layers, mm) + Excel schedule with bar coordinates.
The draft is for the engineer's check and for the draftsman to continue in Revit – not a construction drawing."""
import math, re
import ezdxf
from ezdxf.enums import TextEntityAlignment
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import Circle as MCircle
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import layout as LY

plt.rcParams.update({'font.family': ['DejaVu Sans', 'Loma']})
LAY = {'G-ANNO-TEXT': 7, 'A-DETL-GENF': 7, 'CED_0_10_Solid': 8, 'CED_0_13_Solid': 7, 'CED_0_18_Solid': 7, 'S-DIMS': 3, 'G-ANNO-CHECK': 1}
LW = {'CED_0_10_Solid': 13, 'CED_0_13_Solid': 18, 'CED_0_18_Solid': 35}


class Prims:
    def __init__(self):
        self.p = []

    def L(self, x1, y1, x2, y2, layer='CED_0_13_Solid'):
        self.p.append(('line', layer, (x1, y1, x2, y2)))

    def T(self, x, y, s, h=125, layer='G-ANNO-TEXT', ha='center', rot=0, color=None):
        self.p.append(('text', layer, (x, y, s, h, ha, rot, color)))

    def C(self, x, y, r, layer='A-DETL-GENF', fill=True):
        self.p.append(('dot' if fill else 'circle', layer, (x, y, r)))

    def R(self, x0, y0, x1, y1, layer='CED_0_13_Solid'):
        self.L(x0, y0, x1, y0, layer); self.L(x1, y0, x1, y1, layer); self.L(x1, y1, x0, y1, layer); self.L(x0, y1, x0, y0, layer)


def draw_section(pr, cx, cy, grid, db, counts, T):
    """section centred at (cx, cy); the longer side horizontal. Section axes: x along b, y along h."""
    b, h = grid.b, grid.h
    rot = h >= b                              # drawing X = section y when rotated
    W, H = (h, b) if rot else (b, h)
    x0, y0 = cx - W / 2, cy - H / 2
    def P(x, y):                               # section (x, y) mm -> drawing
        return (x0 + y, y0 + x) if rot else (x0 + x, y0 + y)
    if grid.shape == 'circ':
        pr.C(cx, cy, b / 2, 'CED_0_18_Solid', fill=False)
        pr.C(cx, cy, b / 2 - grid.cover - grid.tie / 2, 'CED_0_13_Solid', fill=False)
    else:
        pr.R(x0, y0, x0 + W, y0 + H, 'CED_0_18_Solid')
        t = grid.cover + grid.tie / 2
        pr.R(x0 + t, y0 + t, x0 + W - t, y0 + H - t)
        rr = db / 2
        for ya, yb_ in T['hoops_h']:          # hoops pairing positions along y (h faces) -> span full b
            (ax, ay), (bx_, by_) = P(t, ya - rr - grid.tie / 2), P(b - t, yb_ + rr + grid.tie / 2)
            pr.R(min(ax, bx_), min(ay, by_), max(ax, bx_), max(ay, by_))
        for xa, xb in T['hoops_b']:
            (ax, ay), (bx_, by_) = P(xa - rr - grid.tie / 2, t), P(xb + rr + grid.tie / 2, h - t)
            pr.R(min(ax, bx_), min(ay, by_), max(ax, bx_), max(ay, by_))
        for yc in T['cross_h']:
            (ax, ay), (bx_, by_) = P(t, yc + rr + grid.tie / 2), P(b - t, yc + rr + grid.tie / 2)
            pr.L(ax, ay, bx_, by_)
        for xc in T['cross_b']:
            (ax, ay), (bx_, by_) = P(xc + rr + grid.tie / 2, t), P(xc + rr + grid.tie / 2, h - t)
            pr.L(ax, ay, bx_, by_)
    for x, y, a in LY.bars_cm(grid, db, counts):
        X, Y = P(x * 10, y * 10)
        pr.C(X, Y, db / 2)
    # dimensions
    pr.L(x0, y0 + H + 220, x0 + W, y0 + H + 220, 'S-DIMS'); pr.L(x0, y0 + H + 120, x0, y0 + H + 300, 'S-DIMS')
    pr.L(x0 + W, y0 + H + 120, x0 + W, y0 + H + 300, 'S-DIMS'); pr.T(cx, y0 + H + 340, f'{W:,.0f}', 110, 'S-DIMS')
    pr.L(x0 - 220, y0, x0 - 220, y0 + H, 'S-DIMS'); pr.L(x0 - 300, y0, x0 - 120, y0, 'S-DIMS')
    pr.L(x0 - 300, y0 + H, x0 - 120, y0 + H, 'S-DIMS'); pr.T(x0 - 330, cy, f'{H:,.0f}', 110, 'S-DIMS', rot=90)
    return W, H


def build(sched, stories, bands, fc_of, project, grids, paths, marks_per_sheet=6):
    """sched: {mark: {band: dict(grid_key, db, counts, nbar, nset, ncross, tie_db, so, s, check, labels)}}"""
    marks = sorted(sched, key=lambda m: int(''.join(c for c in m if c.isdigit()) or 0))
    bands_top = list(reversed(bands))
    maxW = max(max(grids[v['gkey']].b, grids[v['gkey']].h) for m in marks for v in sched[m].values())
    CW_LEV, CW_FC = 2600, 700
    CW_COL = max(3200, maxW + 1400)
    RH = max(2400, min(4200, max(min(grids[v['gkey']].b, grids[v['gkey']].h) for m in marks for v in sched[m].values()) + 2000))
    RH_MARK = 1300
    pr = Prims()
    nrow = len(bands_top)
    TOP = RH_MARK + nrow * RH
    xF1 = CW_LEV + CW_FC
    xEnd = xF1 + CW_COL * len(marks)
    pr.T(xEnd / 2, TOP + 900, 'TABLE OF COLUMN  (DRAFT)', 220)
    pr.T(xEnd / 2, TOP + 550, f"{project.get('name', '')} — draft for engineer check / draftsman", 120)
    pr.R(0, 0, xEnd, TOP, 'CED_0_18_Solid')
    pr.L(CW_LEV, 0, CW_LEV, TOP); pr.L(xF1, 0, xF1, TOP)
    for i in range(len(marks)):
        pr.L(xF1 + CW_COL * i, 0, xF1 + CW_COL * i, TOP)
    pr.L(0, RH_MARK, xEnd, RH_MARK, 'CED_0_18_Solid')
    blist_ = list(reversed(bands_top))                              # bottom -> top, every story range is a row
    for i, (s0, s1) in enumerate(blist_):                          # i = 0 lowest range
        yb = RH_MARK + i * RH; yt = yb + RH
        pr.L(0, yt, CW_LEV, yt, 'CED_0_10_Solid')
        lo = stories[stories.index(s0) - 1] if stories.index(s0) > 0 else 'FDN'
        pr.T(150, yt - 180, f'{s1}', 125, ha='left'); pr.T(150, yb + 180, f'{lo}', 125, ha='left')
        ax_ = CW_LEV * 0.45
        pr.L(ax_, yb + 500, ax_, yt - 500, 'CED_0_10_Solid'); pr.L(ax_, yt - 500, ax_ - 90, yt - 680, 'CED_0_10_Solid')
        pr.L(ax_, yt - 500, ax_ + 90, yt - 680, 'CED_0_10_Solid')
    pr.L(xF1, TOP, xEnd, TOP, 'CED_0_18_Solid')

    def _sig(v):
        return (v['gkey'], v['db'], tuple(v['counts']), v['nset'], v['ncross'], v['tie_db'], v['so'], v['s'])

    def _dc(v):
        m_ = re.search(r'D/C=([0-9.]+)', v['check'][0])
        return float(m_.group(1)) if m_ else 0.0

    # CEDA: every story range is a row; where a mark runs on with the same bars and ties, the cells are merged
    # (no dividing line, section drawn once) – the run is shown continuous
    for mi, m in enumerate(marks):
        x0c, x1c = xF1 + CW_COL * mi, xF1 + CW_COL * (mi + 1)
        i = 0
        while i < len(blist_):
            v = sched[m].get(blist_[i])
            j = i
            if v:
                while j + 1 < len(blist_) and sched[m].get(blist_[j + 1]) and _sig(sched[m][blist_[j + 1]]) == _sig(v):
                    j += 1
            yb, yt = RH_MARK + i * RH, RH_MARK + (j + 1) * RH
            if j + 1 < len(blist_):
                pr.L(x0c, yt, x1c, yt, 'CED_0_10_Solid')
            if v:
                run = [sched[m][blist_[q]] for q in range(i, j + 1)]
                worst = max(run, key=lambda r: (bool(r['check'][1]), _dc(r)))
                g = grids[v['gkey']]
                cx = (x0c + x1c) / 2
                W, H = (max(g.b, g.h), min(g.b, g.h))
                blockH = H + 450 + 260 + 170 * (1 if LY.describe(v['counts'], v.get('phase', 1)) else 0) + 190 + 170 + 190
                cy = (yb + yt) / 2 + blockH / 2 - 450 - H / 2 if j > i else yt - 450 - H / 2
                draw_section(pr, cx, cy, g, v['db'], v['counts'], v['T'])
                yl = cy - H / 2 - 260
                pr.T(cx, yl, f"{v['nbar']} - Ø{v['db']} mm", 125)
                bt = LY.describe(v['counts'], v.get('phase', 1))
                if bt:
                    yl -= 170; pr.T(cx, yl, f'({bt})', 95)
                cr = f" + {v['ncross']} CROSSTIE" if v['ncross'] else ''
                yl -= 190; pr.T(cx, yl, f"{v['nset']}ป{cr} - Ø{v['tie_db']} mm @ {v['so']}  FOR  So", 115)
                yl -= 170; pr.T(cx, yl, f"{v['nset']}ป{cr} - Ø{v['tie_db']} mm @ {v['s']}  FOR  S", 115)
                yl -= 190; pr.T(cx, yl, worst['check'][0], 90, 'G-ANNO-CHECK', color='#c0392b' if worst['check'][1] else '#1f6f3f')
            i = j + 1
    # f'c column (merge consecutive bands with the same f'c)
    i = 0
    blist = list(bands)
    while i < len(blist):
        j = i
        while j + 1 < len(blist) and fc_of[blist[j + 1][0]] == fc_of[blist[i][0]]:
            j += 1
        ya, yb_ = RH_MARK + i * RH, RH_MARK + (j + 1) * RH
        if i > 0:
            pr.L(CW_LEV, ya, xF1, ya, 'CED_0_18_Solid')
        pr.T((CW_LEV + xF1) / 2, (ya + yb_) / 2, f"CONCRETE STRENGTH fc' ( cylinder ) = {fc_of[blist[i][0]]} ksc.", 125, rot=90)
        i = j + 1
    pr.T(300, RH_MARK * 0.68, 'COLUMN', 125, ha='left'); pr.T(CW_LEV - 250, RH_MARK * 0.3, 'NO.', 125, ha='right')
    pr.L(0, RH_MARK, CW_LEV, 0, 'CED_0_10_Solid')
    for mi, m in enumerate(marks):
        cx = xF1 + CW_COL * (mi + 0.5)
        pr.C(cx, RH_MARK * 0.6, 260, 'CED_0_13_Solid', fill=False)
        pr.T(cx, RH_MARK * 0.6, m, 150)
        labs = sorted(set(l for v in sched[m].values() for l in v['labels']), key=lambda s: (len(s), s))
        pr.T(cx, 180, '(' + ', '.join(labs)[:60] + ')', 90, 'G-ANNO-CHECK', color='#555')
    # ---------------------------------------------------------------- DXF
    doc = ezdxf.new('R2013', setup=True); doc.units = ezdxf.units.MM
    doc.styles.add('ARIALUNI', font='ARIALUNI.ttf')
    for n, c in LAY.items():
        lay = doc.layers.add(n, color=c)
        if n in LW:
            lay.dxf.lineweight = LW[n]
    msp = doc.modelspace()
    AL = {'left': TextEntityAlignment.MIDDLE_LEFT, 'center': TextEntityAlignment.MIDDLE_CENTER, 'right': TextEntityAlignment.MIDDLE_RIGHT}
    for kind, layer, d in pr.p:
        if kind == 'line':
            msp.add_line(d[:2], d[2:], dxfattribs={'layer': layer})
        elif kind == 'dot':
            x, y, r_ = d
            msp.add_circle((x, y), r_, dxfattribs={'layer': layer})
            hh = msp.add_hatch(color=7, dxfattribs={'layer': layer}); hh.paths.add_edge_path().add_arc((x, y), r_, 0, 360)
        elif kind == 'circle':
            msp.add_circle(d[:2], d[2], dxfattribs={'layer': layer})
        else:
            x, y, s, hgt, ha, rot, col = d
            t = msp.add_text(s, height=hgt, rotation=rot, dxfattribs={'layer': layer, 'style': 'ARIALUNI'})
            t.set_placement((x, y), align=AL[ha])
    doc.saveas(paths['dxf'])
    # ---------------------------------------------------------------- PDF (A3, marks split over sheets)
    W_all = xEnd
    with PdfPages(paths['pdf']) as pdf:
        for k0 in range(0, len(marks), marks_per_sheet):
            x_lo = 0 if k0 == 0 else xF1 + CW_COL * k0 - (CW_LEV + CW_FC)
            ms = marks[k0:k0 + marks_per_sheet]
            x_hi = xF1 + CW_COL * (k0 + len(ms))
            width = x_hi - x_lo if k0 == 0 else (CW_LEV + CW_FC) + CW_COL * len(ms)
            land = width > TOP * 0.9
            fig = plt.figure(figsize=(16.54, 11.69) if land else (11.69, 16.54))
            ax = fig.add_axes([0.04, 0.04, 0.92, 0.92]); ax.set_aspect('equal'); ax.axis('off')
            xa = 0 if k0 == 0 else xF1 + CW_COL * k0
            ax.set_xlim((-800 if k0 == 0 else xa - 800), x_hi + 800); ax.set_ylim(-400, TOP + 1200)
            fig.canvas.draw()
            bb = ax.get_window_extent()
            sc = min(bb.width / (x_hi + 1600 - (0 if k0 == 0 else xa)), bb.height / (TOP + 1600)) * 72 / fig.dpi
            for kind, layer, d in pr.p:
                if kind == 'line':
                    ax.plot([d[0], d[2]], [d[1], d[3]], color='#2b6cb0' if layer == 'S-DIMS' else 'k',
                            lw={'CED_0_10_Solid': .3, 'CED_0_13_Solid': .5, 'CED_0_18_Solid': 1.0}.get(layer, .3))
                elif kind == 'dot':
                    ax.add_patch(MCircle(d[:2], d[2], color='k'))
                elif kind == 'circle':
                    ax.add_patch(MCircle(d[:2], d[2], fill=False, lw=.7))
                else:
                    x, y, s, hgt, ha, rot, col = d
                    ax.text(x, y, s, fontsize=hgt * sc * 1.35, ha=ha, va='center', rotation=rot,
                            color=col or ('#2b6cb0' if layer == 'S-DIMS' else 'k'), clip_on=True)
            fig.text(0.04, 0.015, 'DRAFT for engineer check — red/green [CHECK] text is on layer G-ANNO-CHECK (freeze before issue).', fontsize=7, color='#555')
            pdf.savefig(fig); plt.close(fig)
    # ---------------------------------------------------------------- Excel schedule
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Schedule'
    hdr = ['Mark', 'ETABS labels', 'Band (dwg level)', 'B (mm)', 'H (mm)', "f'c", 'No. bars', 'Ø', 'Bundles', 'Closed hoops', 'Crossties',
           'Tie Ø', 'So', 'S', 'Lst (mm)', 'Check']
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(name='Arial', bold=True, size=9); c.fill = PatternFill('solid', fgColor='EEF1F4'); c.alignment = Alignment(wrap_text=True)
    ws2 = wb.create_sheet('Bar coordinates')
    ws2.append(['Mark', 'Band', 'Bar', 'x along B (mm)', 'y along H (mm)', 'Ø (mm)'])
    for m in marks:
        for (s0, s1), v in sorted(sched[m].items(), key=lambda kv: -stories.index(kv[0][0])):
            g = grids[v['gkey']]
            lo = stories[stories.index(s0) - 1] if stories.index(s0) > 0 else 'FDN'
            ws.append([m, ', '.join(sorted(v['labels'])), f'{lo} → {s1}', g.b, g.h, fc_of[s0], v['nbar'], v['db'], LY.describe(v['counts'], v.get('phase', 1)) or '-',
                       v['nset'], v['ncross'], v['tie_db'], v['so'], v['s'], v.get('lap') or 'n/a', v['check'][0]])
            for i, (x, y, a) in enumerate(LY.bars_cm(g, v['db'], v['counts']), 1):
                ws2.append([m, f'{lo} → {s1}', i, round(x * 10, 1), round(y * 10, 1), v['db']])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name='Arial', size=9)
    for col_, w in zip('ABCDEFGHIJKLMNOP', [6, 30, 22, 7, 7, 6, 8, 5, 22, 8, 8, 6, 5, 5, 9, 50]):
        ws.column_dimensions[col_].width = w
    wb.save(paths['xlsx'])
