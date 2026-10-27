"""Decision table (Excel) – label × story with my recommendation, f'c decision, mark/band schedule."""
import math
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as CL

F = 'Arial'
f9 = Font(name=F, size=9); fb = Font(name=F, size=9, bold=True); fT = Font(name=F, size=12, bold=True)
fR = Font(name=F, size=9, bold=True, color='B3261E'); fG = Font(name=F, size=9, color='666666', italic=True)
HDR = PatternFill('solid', fgColor='E8EDF2'); RED = PatternFill('solid', fgColor='FBE3E1'); YEL = PatternFill('solid', fgColor='FFF4CC')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def bars_text(r):
    ks = sorted(set(r['counts']))
    t = f"{r['nbar']}-DB{r['db']}"
    if ks != [1]:
        t += ' (bundle ' + '/'.join(str(k) for k in ks if k > 1) + (' at tie corners' if r.get('phase', 0) in (0, 1, 3) else ' at tie corners + other positions') + ')'
    if r.get('phase') == 3:
        t += ' + inner bars at tie crossings'
    return t


def hdr(ws, row, titles):
    for j, t in enumerate(titles, 1):
        c = ws.cell(row, j, t); c.font = fb; c.fill = HDR; c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[row].height = 32


def decision_table(out, path):
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'Summary'
    cfg = out['config']; p = cfg['project']
    ws['A1'] = f"COLUMN DESIGN – DECISION TABLE  ({p.get('name', '')})"; ws['A1'].font = fT
    info = [('Job', p.get('job', '')), ('Seismic', str(cfg.get('seismic', {}))),
            ('Materials', str(cfg['materials'])), ('Detailing', str(cfg.get('detailing', {}))),
            ('Slenderness', str(cfg.get('slender', {}))),
            ("f'c decision", str({k: v for k, v in out['fcinfo'].items() if k != 'table'})),
            ('Bands', ', '.join(f'{a}→{b}' for a, b in out['bands']))]
    for i, (a, b) in enumerate(info, 3):
        ws.cell(i, 1, a).font = fb; ws.cell(i, 2, b).font = f9
    r0 = len(info) + 5
    ws.cell(r0, 1, "f'c study (selected scenario, per story: share of columns at ρ ≤ 1.5 % with typical f'c, failing, under-used D/C < 0.75)").font = fb
    hdr(ws, r0 + 1, ['Story', 'Share ρ ≤ 1.5 %', 'Within the failing limit?', 'Selected', 'Failing', 'Under-used'])
    for i, t in enumerate(out['fcinfo'].get('table', []), r0 + 2):
        ws.cell(i, 1, t['story']).font = f9
        c = ws.cell(i, 2, t['share_low_rho']); c.number_format = '0%'; c.font = f9
        ws.cell(i, 3, 'yes' if t['feasible_above'] else 'NO').font = f9
        if t['story'] == out['fcinfo'].get('reduce_at'):
            ws.cell(i, 4, '← reduce here').font = fR
        if 'fail' in t:
            ws.cell(i, 5, t['fail']).font = f9; ws.cell(i, 6, t['under']).font = f9
    ws.column_dimensions['A'].width = 16; ws.column_dimensions['B'].width = 120
    rn = ws.max_row + 2
    ws.cell(rn, 1, 'Notes').font = fb
    for i, n in enumerate(out['notes'], rn + 1):
        ws.cell(i, 2, n).font = f9
    # ------------------------------------------------------------------ decision rows
    ACC_D = {(str(a), str(b)): float(c) for a, b, c in out.get('config', {}).get('accept_delta', [])}
    wd = wb.create_sheet('Decision')
    hdr(wd, 1, ['Mark', 'ETABS label', 'ETABS story', 'Dwg level (bot→top)', 'Band', 'Section (mm)', "f'c", 'Bars', 'As (cm²)', 'ρ',
                'D/C', 'Governing combo', 'Pu (t)', 'Mux (t-m)', 'Muy (t-m)', 'kLu/r x | y', 'δx | δy max', 'n(δ>1.4)',
                'Shear D/C V2 | V3', 'Ties', 'Status', 'Recommendation'])
    i = 2
    for r in sorted(out['results'], key=lambda r: (r['mark'], r['label'], -out['stories'].index(r['band'][0]))):
        for q in reversed(r['per']):
            rec = []
            if q['status'] == 'ACCEPTED':
                rec.append(f"overload D/C {q['dc']:.3f} accepted by the engineer")
            elif q['status'] == 'REVISE' or q['dc'] > 1:
                rec.append('revise section / f\'c')
            if q['status'] == 'ASK_RHO' or r['As'] / (r['b'] * r['h']) > 0.06:
                rec.append('ρ > 6 %: needs approval')
            if r['As'] / (r['b'] * r['h']) > 0.04:
                rec.append('coupler (ρ > 4 %)')
            if q['n_delta_14']:
                dlim = ACC_D.get((r['label'], q['story']))
                dm = max(q['dmax_x'], q['dmax_y'])
                rec.append(f'δ {dm:.2f} accepted by the engineer (≤ {dlim:.2f})' if dlim and dm <= dlim + 1e-9 else 'δ > 1.4 (ACI 6.2.5.3)')
            sh = max(q['shear']['V2']['dc_lo'], q['shear']['V2']['dc_mid'], q['shear']['V3']['dc_lo'], q['shear']['V3']['dc_mid'])
            if sh > 1:
                rec.append('shear: reduce s / add full-depth legs')
            if not q['ties']['ok']:
                rec.append(f"so > {q['ties']['so_max']:.0f} mm")
            if max(r['counts']) >= 3:
                rec.append('coupler (3-4-bar bundles, CEDA)')
            if r.get('phase') == 3:
                rec.append('inner bars at tie crossings (CEDA pattern, last resort – check)')
            vals = [r['mark'], r['label'], q['story'], f"{q['dwg'][0]}→{q['dwg'][1]}", f"{r['band'][0]}→{r['band'][1]}",
                    f"{r['b']*10:.0f}×{r['h']*10:.0f}" if r['shape'] == 'rect' else f"Ø{r['b']*10:.0f}", q['fc'], bars_text(r),
                    round(r['As'], 1), r['As'] / (r['b'] * r['h']) if r['shape'] == 'rect' else r['As'] / (math.pi * r['b'] ** 2 / 4),
                    round(q['dc'], 3), q['gov'].get('combo', ''), round(q['gov']['P'] / 1000, 1), round(q['gov']['Mx'] / 1000, 1),
                    round(q['gov']['My'] / 1000, 1), f"{q['slr_x']:.1f} | {q['slr_y']:.1f}", f"{q['dmax_x']:.2f} | {q['dmax_y']:.2f}",
                    q['n_delta_14'], f"{max(q['shear']['V2']['dc_lo'], q['shear']['V2']['dc_mid']):.2f} | "
                    f"{max(q['shear']['V3']['dc_lo'], q['shear']['V3']['dc_mid']):.2f}",
                    f"{r['nset']}ป{(' +' + str(r['ncross']) + ' x-tie') if r['ncross'] else ''} DB{r['tie_db']} @{r['so']}/{r['s']}",
                    q['status'], '; '.join(rec) or 'O.K.']
            for j, v in enumerate(vals, 1):
                c = wd.cell(i, j, v); c.font = f9; c.border = BOX
            wd.cell(i, 10).number_format = '0.00%'
            if q['status'] == 'REVISE' or (q['dc'] > 1 and q['status'] != 'ACCEPTED'):
                for j in range(1, len(vals) + 1): wd.cell(i, j).fill = RED
            elif q['status'] in ('ASK_RHO', 'ACCEPTED'):
                for j in range(1, len(vals) + 1): wd.cell(i, j).fill = YEL
            i += 1
    for j, w in enumerate([6, 9, 9, 12, 12, 11, 6, 30, 8, 7, 7, 34, 8, 9, 9, 11, 11, 7, 12, 22, 9, 45], 1):
        wd.column_dimensions[CL(j)].width = w
    wd.freeze_panes = 'D2'
    # ------------------------------------------------------------------ schedule (mark × band)
    wsch = wb.create_sheet('Schedule')
    hdr(wsch, 1, ['Mark', 'ETABS labels', 'Band (bot→top)', 'Section', "f'c", 'Bars', 'Ties', 'Lap Lst (mm)', 'Max D/C', 'Grid notes'])
    i = 2
    seen = set()
    for r in sorted(out['results'], key=lambda r: (r['mark'], -out['stories'].index(r['band'][0]))):
        k = (r['mark'], tuple(r['band']))
        if k in seen:
            continue
        seen.add(k)
        labs = sorted(l for l, m in out['marks'].items() if m == r['mark'])
        dcm = max(q['dc'] for rr in out['results'] if rr['mark'] == r['mark'] and rr['band'] == r['band'] for q in rr['per'])
        vals = [r['mark'], ', '.join(labs), f"{r['band'][0]}→{r['band'][1]}", f"{r['b']*10:.0f}×{r['h']*10:.0f}", r['fc'], bars_text(r),
                f"{r['nset']}ป DB{r['tie_db']} @{r['so']} (lo) / @{r['s']}", r['lap'] or 'n/a – add to lap table', round(dcm, 3),
                '; '.join(r['grid_notes'])]
        for j, v in enumerate(vals, 1):
            c = wsch.cell(i, j, v); c.font = f9; c.border = BOX
        i += 1
    for j, w in enumerate([6, 30, 14, 11, 6, 34, 26, 14, 8, 60], 1):
        wsch.column_dimensions[CL(j)].width = w
    wsch.cell(i + 1, 1, 'Bundles: started at the tie corners, each bundle size completed over the section before the next (CEDA 20 Oct 2026); 3-4-bar bundles and ρ > 4 %: couplers.').font = fG
    wb.save(path)
